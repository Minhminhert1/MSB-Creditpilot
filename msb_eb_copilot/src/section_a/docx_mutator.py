"""Authoritative MB07 DOCX Section A Mutator (Phases 12 & 13).

Provides deterministic in-place OOXML mutations for Table 1 (Section A) of the
authoritative MB07 template, handling:
- Standard and merged cell text updates
- Prefixed and suffixed label text
- Legacy FORMCHECKBOX toggle states
- Content control (w:sdt) dropdown selection text
- Stale sample/placeholder cleanup
"""

from __future__ import annotations

import os
from typing import Any, Sequence
import docx
from docx.oxml.ns import qn
from docx.table import _Cell, Table
from lxml import etree


from ..template_rendering.safe_mutation import SafeCellMutator


class SectionADocxMutator:
    """Mutates Section A (Table 1) of an authoritative MB07 Word document."""

    def __init__(self, doc_or_path: docx.Document | str, table_index: int = 1) -> None:
        if isinstance(doc_or_path, str):
            if not os.path.exists(doc_or_path):
                raise FileNotFoundError(f"DOCX file not found at '{doc_or_path}'.")
            self.document = docx.Document(doc_or_path)
            self.file_path: str | None = doc_or_path
        else:
            self.document = doc_or_path
            self.file_path = None

        if len(self.document.tables) <= table_index:
            raise ValueError(
                f"Document does not contain table at index {table_index}. "
                f"Total tables: {len(self.document.tables)}"
            )
        self.table: Table = self.document.tables[table_index]
        self._table_index = table_index
        # Determine row offset between old template (30 rows) and new template (32 rows with DNTN)
        self.row_offset = 2 if (len(self.table.rows) >= 32 or (len(self.table.rows) > 0 and "chủ DNTN" in self.table.rows[0].cells[0].text)) else 0

    def get_unique_cells(self, row_index: int) -> list[_Cell]:
        """Return the deduplicated, ordered cells for a given row index.

        Because python-docx repeats merged cells across columns, deduplication
        preserves the actual physical XML cells in logical left-to-right order.
        """
        if row_index < 0 or row_index >= len(self.table.rows):
            raise IndexError(
                f"Row index {row_index} out of range [0, {len(self.table.rows) - 1}]."
            )
        row = self.table.rows[row_index]
        return list(dict.fromkeys(row.cells))

    def set_cell_text(self, row_index: int, cell_index: int, text: str) -> None:
        """Set the text of a specific physical cell in Section A using run-preserving SafeCellMutator."""
        unique_cells = self.get_unique_cells(row_index)
        if cell_index < 0 or cell_index >= len(unique_cells):
            raise IndexError(
                f"Cell index {cell_index} out of range [0, {len(unique_cells) - 1}] "
                f"for row {row_index}."
            )
        cell = unique_cells[cell_index]
        SafeCellMutator.set_cell_text(cell, text)

    def set_prefixed_text(
        self,
        row_index: int,
        cell_index: int,
        prefix: str | None = None,
        value: str | None = None,
        suffix: str | None = None,
    ) -> None:
        """Set cell text with optional prefix and suffix formatting."""
        parts: list[str] = []
        if prefix and prefix.strip():
            parts.append(prefix.strip())
        if value is not None and str(value).strip():
            parts.append(str(value).strip())
        if suffix and suffix.strip():
            parts.append(suffix.strip())

        formatted = " ".join(parts)
        self.set_cell_text(row_index, cell_index, formatted)

    def set_checkbox(self, row_index: int, selected_index: int | None, selected_value: str | None = None) -> None:
        """Toggle legacy FORMCHECKBOX controls or modern Content Control dropdowns in a row.

        The checkbox at selected_index (0-based) is marked checked (w:default=1,
        w:checked=1), and all other checkboxes in the row are unchecked.
        If selected_index is None, all checkboxes in the row are unchecked.
        If the row uses w:sdt (Content Control dropdown) instead of FORMCHECKBOX,
        the corresponding item is selected and written to the display text.
        """
        if row_index < 0 or row_index >= len(self.table.rows):
            raise IndexError(
                f"Row index {row_index} out of range [0, {len(self.table.rows) - 1}]."
            )
        row = self.table.rows[row_index]
        ffdata_elements = row._tr.xpath(".//w:ffData")

        if ffdata_elements:
            for idx, ff in enumerate(ffdata_elements):
                cb = ff.find(qn("w:checkBox"))
                if cb is None:
                    continue

                is_checked = (selected_index is not None and idx == selected_index)
                val = "1" if is_checked else "0"

                # Set default
                df = cb.find(qn("w:default"))
                if df is None:
                    df = etree.SubElement(cb, qn("w:default"))
                df.set(qn("w:val"), val)

                # Set checked
                chk = cb.find(qn("w:checked"))
                if chk is None:
                    chk = etree.SubElement(cb, qn("w:checked"))
                chk.set(qn("w:val"), val)
        else:
            # Check for modern w:sdt content control
            sdt_elements = row._tr.xpath(".//w:sdt")
            if sdt_elements:
                sdt = sdt_elements[0]
                items = sdt.xpath(".//w:listItem/@w:value", namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"})
                filtered = [it for it in items if it not in ("Choose an item.", "Tick chọn", "Chọn kết quả")]
                target_text = None
                if selected_value:
                    val_clean = selected_value.strip().lower()
                    # Custom alias maps for canonical values to dropdown text
                    alias_map = {
                        "kh_moi": "kh mới",
                        "kh_hien_huu": "kh hiện hữu",
                        "co": "có",
                        "khong": "không",
                        "bat_buoc_danh_gia": "thuộc đối tượng phải đánh giá rủi ro môi trường xã hội",
                        "khong_bat_buoc_danh_gia": "không thuộc đối tượng phải đánh giá rủi ro môi trường xã hội",
                        "trong_gioi_han": "trong giới hạn",
                        "vuot_gioi_han": "vượt giới hạn",
                        "cap_moi": "cấp mới",
                        "tai_cap": "tái cấp",
                    }
                    normalized_target = alias_map.get(val_clean, val_clean)
                    for it in filtered:
                        it_lower = it.strip().lower()
                        if it_lower == normalized_target or normalized_target in it_lower or it_lower in normalized_target:
                            target_text = it
                            break
                if not target_text and selected_index is not None and 0 <= selected_index < len(filtered):
                    target_text = filtered[selected_index]
                if target_text:
                    self.set_sdt_dropdown_text(row_index, target_text)

    def set_sdt_dropdown_text(self, row_index: int, text: str) -> None:
        """Set the display text for a row-level content control (w:sdt).

        Updates the text inside w:sdtContent and removes placeholder color/italic styling.
        """
        if row_index < 0 or row_index >= len(self.table.rows):
            raise IndexError(
                f"Row index {row_index} out of range [0, {len(self.table.rows) - 1}]."
            )
        row = self.table.rows[row_index]
        sdt_contents = row._tr.xpath(".//w:sdtContent")
        if not sdt_contents:
            raise ValueError(f"No w:sdtContent found in row {row_index}.")

        sdt_content = sdt_contents[0]
        t_nodes = sdt_content.findall(".//" + qn("w:t"))
        if not t_nodes:
            raise ValueError(f"No w:t text nodes found in w:sdtContent of row {row_index}.")

        t_nodes[0].text = text
        for extra_t in t_nodes[1:]:
            extra_t.text = ""

        # Remove red color and italics placeholder styling
        for color in sdt_content.findall(".//" + qn("w:color")):
            parent = color.getparent()
            if parent is not None:
                parent.remove(color)
        for i_tag in sdt_content.findall(".//" + qn("w:i")):
            parent = i_tag.getparent()
            if parent is not None:
                parent.remove(i_tag)

    def clear_sample_placeholders(self) -> None:
        """Clear known stale sample and placeholder values from Table 1."""
        offset = self.row_offset
        # Company name
        self.set_cell_text(0 + offset, 1, "")

        # SDT placeholder
        try:
            self.set_sdt_dropdown_text(2 + offset, "")
        except Exception:
            pass

        # Sample industry code, name, share
        self.set_cell_text(15 + offset, 1, "")
        self.set_cell_text(15 + offset, 2, "")
        self.set_cell_text(15 + offset, 3, "")

        # Sample main products
        self.set_cell_text(16 + offset, 2, "")

        # Sample total trđ
        self.set_cell_text(26 + offset, 2, "")
        self.set_cell_text(26 + offset, 3, "")

    def save(self, target_path: str) -> str:
        """Save the mutated document to a file path."""
        os.makedirs(os.path.dirname(os.path.abspath(target_path)), exist_ok=True)
        self.document.save(target_path)
        return target_path
