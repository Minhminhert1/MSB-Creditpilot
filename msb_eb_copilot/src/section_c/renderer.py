"""Module: renderer.py
Mô tả: Engine render Phần C vào văn bản Word MB07 (Bảng biểu chuẩn Bảng 01 - Bảng 06 của MSB).
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from .models import SectionCData


class SectionCRenderer:
    """Renderer Phần C cho Tờ trình MB07."""

    @staticmethod
    def _set_cell_background(cell, fill_hex: str = "F1F5F9"):
        tcPr = cell._element.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    @staticmethod
    def _set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
        tcPr = cell._element.get_or_add_tcPr()
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top}" w:type="dxa"/>'
            f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
            f'<w:left w:w="{left}" w:type="dxa"/>'
            f'<w:right w:w="{right}" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)

    @staticmethod
    def _set_table_borders(table, color="B0C4DE", sz="4"):
        tblPr = table._element.xpath('w:tblPr')
        if tblPr:
            borders = parse_xml(
                f'<w:tblBorders {nsdecls("w")}>'
                f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                f'<w:left w:val="none"/>'
                f'<w:right w:val="none"/>'
                f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
                f'<w:insideV w:val="none"/>'
                f'</w:tblBorders>'
            )
            tblPr[0].append(borders)

    @classmethod
    def _create_styled_table(cls, doc, headers, rows, col_widths=None, align_right_cols=None):
        align_right_cols = align_right_cols or []
        table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cls._set_table_borders(table)

        # Format Headers
        for col_idx, text in enumerate(headers):
            cell = table.rows[0].cells[col_idx]
            cell.text = text
            cls._set_cell_background(cell, "E2E8F0")
            cls._set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if col_idx in align_right_cols else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                run.bold = True
                run.font.name = "Times New Roman"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0, 32, 96)

        # Format Data Rows
        for row_idx, r_data in enumerate(rows):
            row = table.rows[row_idx + 1]
            bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
            for col_idx, val in enumerate(r_data):
                cell = row.cells[col_idx]
                cell.text = str(val) if val is not None else ""
                cls._set_cell_background(cell, bg)
                cls._set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if col_idx in align_right_cols else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9.0)
                    run.font.color.rgb = RGBColor(51, 65, 85)

        if col_widths:
            for i, w in enumerate(col_widths):
                for row in table.rows:
                    row.cells[i].width = Inches(w)

        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(2)
        sp.paragraph_format.space_after = Pt(4)
        return table

    @classmethod
    def render_to_document(cls, doc: docx.Document, data: SectionCData):
        """Render nội dung Phần C vào đối tượng Document."""
        def add_h2(text):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(text)
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(0, 32, 96)
            return p

        def add_body(text, bold_prefix=None):
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(3)
            if bold_prefix:
                rb = p.add_run(bold_prefix)
                rb.bold = True
                rb.font.name = "Times New Roman"
                rb.font.size = Pt(10)
            rt = p.add_run(text)
            rt.font.name = "Times New Roman"
            rt.font.size = Pt(10)
            return p

        # 1. Thời gian hoạt động & Quá trình tăng vốn
        add_h2("1. Thời gian hoạt động của doanh nghiệp")
        add_body(data.history_narrative, bold_prefix="Tóm tắt quá trình hình thành, phát triển: ")

        if data.capital_milestones:
            headers_m = ["Thời điểm", "Vốn điều lệ (triệu VND)", "Sự kiện ghi nhận"]
            rows_m = [
                [m.effective_date, f"{m.charter_capital_million_vnd:,.0f}".replace(",", "."), m.event_description]
                for m in data.capital_milestones
            ]
            cls._create_styled_table(doc, headers_m, rows_m, col_widths=[1.5, 2.0, 3.2], align_right_cols=[1])

        # 2. Chủ sở hữu & Ban lãnh đạo
        add_h2("2. Chủ sở hữu, Người có liên quan & Ban lãnh đạo")
        if data.parent_company_or_owner:
            add_body(f"Doanh nghiệp thuộc sở hữu của {data.parent_company_or_owner}.", bold_prefix="Thông tin Chủ sở hữu: ")

        if data.major_shareholders:
            headers_sh = ["STT", "Tên Cổ đông / Thành viên", "Mã số thuế / CCCD", "Tỷ lệ sở hữu (%)", "Vốn góp (triệu VND)"]
            rows_sh = [
                [sh.stt, sh.shareholder_name, sh.id_tax_code, f"{sh.ownership_percentage:.2f}%", f"{sh.contributed_capital_million_vnd:,.0f}".replace(",", ".")]
                for sh in data.major_shareholders
            ]
            cls._create_styled_table(doc, headers_sh, rows_sh, col_widths=[0.6, 2.5, 1.4, 1.2, 1.3], align_right_cols=[3, 4])

        bl_text = "Hợp lệ, không vi phạm Điều 126 Luật TCTD, không thuộc danh sách Blacklist." if data.blacklist_status.value == "KHONG_VI_PHAM" else "CẢNH BÁO: Vi phạm danh sách đen."
        add_body(bl_text, bold_prefix="Xác nhận tình trạng pháp lý: ")

        if data.management_members:
            for mb in data.management_members:
                add_body(f"{mb.full_name} ({mb.years_at_company} năm tại DN) — {mb.profile_summary}", bold_prefix=f"+ {mb.position}: ")

        # 3. Sản phẩm, Quy trình sản xuất & Cơ sở vật chất
        add_h2("3. Ngành nghề, Sản phẩm, Cơ sở vật chất và MMTB")
        if data.products:
            headers_p = ["STT", "Sản phẩm chính", "Thương hiệu", "Tỷ trọng Doanh thu (%)"]
            rows_p = [
                [p.stt, p.product_name, p.brand_name, f"{p.revenue_share_percentage:.1f}%"]
                for p in data.products
            ]
            cls._create_styled_table(doc, headers_p, rows_p, col_widths=[0.6, 2.8, 1.8, 1.8], align_right_cols=[3])

        if data.production_technology_summary:
            add_body(data.production_technology_summary, bold_prefix="Công nghệ & Quy trình sản xuất: ")

        # Bảng 01: Kho bãi
        if data.warehouses:
            add_body("Bảng 01: Năng lực lưu trữ hàng hóa (Hệ thống nhà xưởng, kho bãi):", bold_prefix="")
            headers_wh = ["STT", "Loại cơ sở", "Địa chỉ", "Diện tích (m²)", "Hình thức", "Công suất lưu trữ"]
            rows_wh = [
                [w.stt, w.facility_type, w.address, f"{w.area_m2:,.0f}", w.ownership_type, w.capacity_description]
                for w in data.warehouses
            ]
            cls._create_styled_table(doc, headers_wh, rows_wh, col_widths=[0.5, 1.2, 2.2, 1.0, 1.0, 1.1], align_right_cols=[3])

        # Bảng 02: MMTB
        if data.equipments:
            add_body("Bảng 02: Chi tiết hệ thống Máy móc thiết bị (MMTB):", bold_prefix="")
            headers_eq = ["STT", "Tên thiết bị / Dây chuyền", "Xuất xứ & Công nghệ", "Công suất thiết kế", "Hiệu suất vận hành"]
            rows_eq = [
                [eq.stt, eq.equipment_name, eq.origin_and_technology, eq.designed_capacity, eq.utilization_rate]
                for eq in data.equipments
            ]
            cls._create_styled_table(doc, headers_eq, rows_eq, col_widths=[0.5, 2.3, 1.8, 1.4, 1.0], align_right_cols=[4])

        # 4. Thị trường Đầu vào & Top Nhà cung cấp (Bảng 04)
        add_h2("4. Thị trường Đầu vào và Nhà cung cấp chính")
        if data.raw_materials_overview:
            add_body(data.raw_materials_overview, bold_prefix="Nguyên vật liệu chính: ")

        if data.suppliers:
            add_body("Bảng 04: Danh sách Nhà cung cấp chính:", bold_prefix="")
            headers_sup = ["STT", "Tên Nhà cung cấp", "Mặt hàng cung ứng", "Tỷ trọng mua (%)", "Phương thức thanh toán"]
            rows_sup = [
                [s.stt, s.supplier_name, s.supplied_goods, f"{s.purchase_share_percentage:.1f}%", s.payment_terms]
                for s in data.suppliers
            ]
            cls._create_styled_table(doc, headers_sup, rows_sup, col_widths=[0.5, 2.5, 1.8, 1.0, 1.4], align_right_cols=[3])

        # 5. Thị trường Đầu ra & Top Khách hàng (Bảng 06)
        add_h2("5. Thị trường Đầu ra, Khách hàng và Đối thủ cạnh tranh")
        if data.distribution_channels:
            add_body(data.distribution_channels, bold_prefix="Kênh phân phối: ")

        comp_txt = f"Thị phần ước tính: {data.market_share_estimate}. Đối thủ chính: {', '.join(data.top_competitors)}. Lợi thế cạnh tranh: {data.competitive_advantages}"
        add_body(comp_txt, bold_prefix="Thị phần & Cạnh tranh: ")

        if data.customers:
            add_body("Bảng 06: Danh sách Khách hàng đầu ra chính:", bold_prefix="")
            headers_cust = ["STT", "Tên Khách hàng", "Sản phẩm cung cấp", "Tỷ trọng Doanh thu (%)", "Chính sách công nợ"]
            rows_cust = [
                [c.stt, c.customer_name, c.product_purchased, f"{c.revenue_share_percentage:.1f}%", c.credit_terms]
                for c in data.customers
            ]
            cls._create_styled_table(doc, headers_cust, rows_cust, col_widths=[0.5, 2.5, 1.8, 1.2, 1.2], align_right_cols=[3])

        # 6. Nhận xét & Đánh giá của RM
        if data.rm_supply_chain_assessment or data.rm_credit_risk_mitigation:
            add_h2("6. Nhận xét, đánh giá của ĐVKD về Chuỗi cung ứng & Hoạt động kinh doanh")
            if data.rm_supply_chain_assessment:
                add_body(data.rm_supply_chain_assessment, bold_prefix="+ Đánh giá chuỗi cung ứng & sự ổn định: ")
            if data.rm_credit_risk_mitigation:
                add_body(data.rm_credit_risk_mitigation, bold_prefix="+ Biện pháp kiểm soát rủi ro & điều kiện tín dụng: ")

    @classmethod
    def generate_docx(cls, data: SectionCData, output_path: str) -> str:
        """Tạo file Word độc lập cho Phần C."""
        doc = docx.Document()
        for s in doc.sections:
            s.top_margin = Inches(0.7)
            s.bottom_margin = Inches(0.7)
            s.left_margin = Inches(0.75)
            s.right_margin = Inches(0.75)

        # Header Title
        p_title = doc.add_paragraph()
        r_title = p_title.add_run("PHẦN C: HOẠT ĐỘNG KINH DOANH CỦA KHÁCH HÀNG")
        r_title.bold = True
        r_title.font.name = "Times New Roman"
        r_title.font.size = Pt(12)
        r_title.font.color.rgb = RGBColor(0, 32, 96)

        cls.render_to_document(doc, data)

        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        doc.save(output_path)
        return output_path
