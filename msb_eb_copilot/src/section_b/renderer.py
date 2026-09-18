# -*- coding: utf-8 -*-
"""Section B Pipeline Orchestrator & High-Fidelity DOCX Renderer."""

import os
import sys
from typing import Dict, Any, Tuple, Optional, List
import docx

from .validator import validate_and_calculate_section_b
from .docx_mutator import (
    remove_unselected_blocks,
    mutate_section_1_table,
    populate_detail_tables,
    update_opening_paragraph
)
from .models import SectionBDerivedTotals, SectionBValidationReport


def resolve_template_path(custom_template_path: Optional[str] = None) -> str:
    """Tìm đường dẫn template 'PHẦN B.docx' hoặc 'PHAN_B_TEMPLATE.docx' tin cậy."""
    if custom_template_path and os.path.exists(custom_template_path):
        return custom_template_path

    # Candidate paths
    candidates = [
        "PHẦN B.docx",
        os.path.join("msb_eb_copilot", "templates", "PHẦN B.docx"),
        os.path.join("msb_eb_copilot", "templates", "PHAN_B_TEMPLATE.docx"),
        os.path.abspath("PHẦN B.docx"),
        os.path.abspath(os.path.join("msb_eb_copilot", "templates", "PHAN_B_TEMPLATE.docx")),
        os.path.abspath(os.path.join("..", "templates", "PHAN_B_TEMPLATE.docx")),
    ]

    for p in candidates:
        if os.path.exists(p):
            return p

    raise FileNotFoundError(
        "Không tìm thấy file template 'PHẦN B.docx' hoặc 'PHAN_B_TEMPLATE.docx'. "
        "Vui lòng kiểm tra lại thư mục hoặc cung cấp đường dẫn hợp lệ."
    )


def render_section_b_docx(
    input_data: Dict[str, Any],
    output_path: str,
    template_path: Optional[str] = None
) -> Tuple[str, SectionBDerivedTotals, SectionBValidationReport]:
    """Quy trình sinh tài liệu Phần B hoàn chỉnh trực tiếp từ template gốc.
    
    Returns:
        output_path: Đường dẫn file Word đã sinh
        totals: Dữ liệu tính toán tổng hợp
        report: Báo cáo validation (warnings, blocking errors)
    """
    # 1. Validation & Tính toán tự động
    totals, report, processed_facs = validate_and_calculate_section_b(input_data)
    if not report.is_valid:
        raise ValueError(f"Dữ liệu đầu vào không hợp lệ: {report.blocking_errors}")

    # 2. Định vị template gốc
    tmpl = resolve_template_path(template_path)
    doc = docx.Document(tmpl)

    raw_needs = input_data.get("selected_needs", [])
    selected_needs = [n.value if hasattr(n, "value") else str(n) for n in raw_needs]
    currency = input_data.get("currency", "VND")

    # 3. Cập nhật đoạn văn bản mở đầu Phần B (P2)
    update_opening_paragraph(doc, totals)

    # 4. Điền dữ liệu vào các Bảng chi tiết Mục 2
    populate_detail_tables(doc, selected_needs, processed_facs, currency=currency)

    # 5. Xử lý Bảng tổng hợp Mục 1 (Table 0): Xóa dòng không áp dụng & cập nhật checkbox
    mutate_section_1_table(doc, selected_needs, processed_facs, totals)

    # 6. Xóa các Block 2.x không được chọn khỏi body
    remove_unselected_blocks(doc, selected_needs)

    # 7. Lưu file kết quả
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc.save(output_path)

    # 8. Post-generation Document QA
    _verify_generated_document(output_path, selected_needs)

    return output_path, totals, report


def _verify_generated_document(output_path: str, selected_needs: List[str]):
    """QA kiểm tra tính toàn vẹn của file DOCX vừa sinh."""
    if not os.path.exists(output_path) or os.path.getsize(output_path) < 1000:
        raise RuntimeError(f"Lỗi: File {output_path} rỗng hoặc không tồn tại.")

    doc = docx.Document(output_path)
    # Verify no red placeholder remains in detail tables
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        if r.font.color and r.font.color.rgb:
                            hex_color = str(r.font.color.rgb)
                            if hex_color.upper() == "FF0000" and "ghi số tiền" in r.text.lower():
                                raise AssertionError("Lỗi QA: Vẫn còn placeholder màu đỏ trong bảng chi tiết!")
