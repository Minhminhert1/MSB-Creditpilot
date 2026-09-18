from __future__ import annotations
import os
from typing import List, Optional, Union
import docx

from ..template_verification.structure_fingerprint import (
    DocumentStructureSnapshot,
    DynamicTableRule,
    StructureComparisonResult,
    compare_template_structure,
    snapshot_template_structure,
)


class MB07FidelityError(Exception):
    """Raised when output document violates the MB07 structural fidelity contract."""
    pass


class UnsupportedTemplateVersionError(Exception):
    """Raised when an input DOCX template is unrecognized or structurally incompatible."""
    pass


def validate_template_compatibility(doc_or_path: Union[docx.Document, str]) -> bool:
    """Validate whether a document conforms to the supported MB07 template structure.
    
    Checks for essential section anchors, headings, and minimum structural elements.
    Raises UnsupportedTemplateVersionError if incompatible.
    """
    if isinstance(doc_or_path, str):
        if not os.path.exists(doc_or_path):
            raise FileNotFoundError(f"Template file not found at '{doc_or_path}'")
        doc = docx.Document(doc_or_path)
    else:
        doc = doc_or_path

    # Must contain at least 20 tables in standard full MB07
    if len(doc.tables) < 20:
        raise UnsupportedTemplateVersionError(
            f"Incompatible template: expected at least 20 tables for full MB07, found {len(doc.tables)}."
        )

    # Search for canonical anchors across paragraphs and tables
    all_p_text = " ".join(p.text for p in doc.paragraphs)
    all_tbl_text = " ".join(c.text for tbl in doc.tables for r in tbl.rows for c in r.cells)
    combined_text = (all_p_text + " " + all_tbl_text).lower()

    required_anchors = [
        "đơn vị trình",
        "tên khách hàng",
        "nội dung đề xuất cấp tín dụng",
        "hoạt động kinh doanh của khách hàng",
        "tình hình tài chính doanh nghiệp",
        "thông tin quan hệ tín dụng của khách hàng",
    ]

    missing = [a for a in required_anchors if a not in combined_text]
    if missing:
        raise UnsupportedTemplateVersionError(
            f"Incompatible template: missing required semantic anchors: {', '.join(missing)}"
        )

    return True


class StructureGuard:
    """Quality Gate Guard that enforces MB07 template preservation."""

    def __init__(
        self,
        template_path: str,
        dynamic_table_rules: Optional[List[DynamicTableRule]] = None,
        allow_table_expansion: bool = False,
        allow_row_growth: bool = False,
    ):
        self.template_path = template_path
        self.dynamic_table_rules = dynamic_table_rules or []
        self.allow_table_expansion = allow_table_expansion
        self.allow_row_growth = allow_row_growth

        # Validate baseline template compatibility
        validate_template_compatibility(template_path)

        # Capture baseline structural snapshot
        self.baseline_snapshot: DocumentStructureSnapshot = snapshot_template_structure(template_path)

    def verify(self, output_docx_path: str, raise_on_violation: bool = True) -> StructureComparisonResult:
        """Verify output document structure against the baseline snapshot."""
        if not os.path.exists(output_docx_path):
            raise FileNotFoundError(f"Output document not found at '{output_docx_path}'")

        output_snapshot = snapshot_template_structure(output_docx_path)

        result = compare_template_structure(
            template_snapshot=self.baseline_snapshot,
            output_snapshot=output_snapshot,
            allow_table_expansion=self.allow_table_expansion,
            allow_row_growth=self.allow_row_growth,
            dynamic_table_rules=self.dynamic_table_rules,
        )

        if not result.is_structurally_sound and raise_on_violation:
            violation_summary = "\n- " + "\n- ".join(result.all_violations)
            raise MB07FidelityError(
                f"MB07 TEMPLATE FIDELITY FAILED: Document structure violated {len(result.all_violations)} invariants:{violation_summary}"
            )

        return result
