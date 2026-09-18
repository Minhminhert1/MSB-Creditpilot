"""Module: renderer.py
Mô tả: Engine render Phần E (Quan hệ tín dụng & Báo cáo CIC - Bảng 07) vào văn bản Word MB07.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from .models import SectionEData


class SectionERenderer:
    """Renderer Phần E chuẩn format Bảng 07 của MSB."""

    @staticmethod
    def _set_cell_background(cell, fill_hex: str = "F1F5F9"):
        tcPr = cell._element.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    @staticmethod
    def _set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
        tcPr = cell._element.get_or_add_tcPr()
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top}" w:type="dxa"/>'
            f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
            f'<w:left w:w="{left}" w:type="dxa"/>'
            f'<w:right w:w="{right}" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)

    @staticmethod
    def _set_table_borders(table, color="B0C4DE", sz="4"):
        tblPr = table._element.xpath('w:tblPr')
        if tblPr:
            borders = parse_xml(
                f'<w:tblBorders {nsdecls("w")}>'
                f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                f'<w:left w:val="none"/>'
                f'<w:right w:val="none"/>'
                f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
                f'<w:insideV w:val="none"/>'
                f'</w:tblBorders>'
            )
            tblPr[0].append(borders)

    @classmethod
    def _create_styled_table(cls, doc, headers, rows, col_widths=None, align_right_cols=None):
        align_right_cols = align_right_cols or []
        table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cls._set_table_borders(table)

        # Headers
        for col_idx, text in enumerate(headers):
            cell = table.rows[0].cells[col_idx]
            cell.text = text
            cls._set_cell_background(cell, "E2E8F0")
            cls._set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if col_idx in align_right_cols else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                run.bold = True
                run.font.name = "Times New Roman"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0, 32, 96)

        # Data Rows
        for row_idx, r_data in enumerate(rows):
            row = table.rows[row_idx + 1]
            is_total_row = (row_idx == len(rows) - 1) and ("Tổng" in str(r_data[1]))
            bg = "E2E8F0" if is_total_row else ("F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
            for col_idx, val in enumerate(r_data):
                cell = row.cells[col_idx]
                cell.text = str(val) if val is not None else ""
                cls._set_cell_background(cell, bg)
                cls._set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if col_idx in align_right_cols else (WD_ALIGN_PARAGRAPH.CENTER if col_idx == 0 or col_idx == 8 else WD_ALIGN_PARAGRAPH.LEFT)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9.0)
                    if is_total_row:
                        run.bold = True
                        run.font.color.rgb = RGBColor(0, 32, 96)
                    else:
                        run.font.color.rgb = RGBColor(51, 65, 85)

        if col_widths:
            for i, w in enumerate(col_widths):
                for row in table.rows:
                    row.cells[i].width = Inches(w)

        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(2)
        sp.paragraph_format.space_after = Pt(4)
        return table

    @classmethod
    def render_to_document(cls, doc: docx.Document, data: SectionEData):
        """Render nội dung Phần E vào đối tượng Document."""
        def add_h2(text):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(text)
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(0, 32, 96)
            return p

        def add_body(text, bold_prefix=None):
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(3)
            if bold_prefix:
                rb = p.add_run(bold_prefix)
                rb.bold = True
                rb.font.name = "Times New Roman"
                rb.font.size = Pt(10)
            rt = p.add_run(text)
            rt.font.name = "Times New Roman"
            rt.font.size = Pt(10)
            return p

        def fmt_curr(v):
            if v is None:
                return "-"
            if v == 0:
                return "0"
            try:
                f = float(v)
                return f"{f:,.0f}".replace(",", ".")
            except (ValueError, TypeError):
                return str(v)

        # Tiêu đề mục
        add_h2("1. Quan hệ tín dụng tại các TCTD, Định chế tài chính (Theo CIC)")
        add_body(f"Ngày tra cứu CIC: {data.cic_report_date}. Khách hàng hiện có quan hệ tín dụng với {data.total_institutions_count} TCTD.")

        # Bảng 07: Thông tin quan hệ tín dụng
        p_dvt = doc.add_paragraph()
        p_dvt.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_dvt.paragraph_format.space_after = Pt(2)
        r_dvt = p_dvt.add_run("ĐVT: Triệu đồng")
        r_dvt.italic = True
        r_dvt.font.name = "Times New Roman"
        r_dvt.font.size = Pt(9.5)

        headers_e = [
            "TT", "Tên TCTD", "Tổng HMTD ngắn hạn", 
            "Dư nợ NH (VND)", "Dư nợ NH (USD)", "Dư nợ TDH", 
            "Tổng dư nợ", "Tài sản bảo đảm", "Nhóm nợ"
        ]

        rows_e = []
        tot_limit = 0.0
        tot_vnd = 0.0
        tot_usd = 0.0
        tot_tdh = 0.0
        tot_debt = 0.0

        for r in data.relations:
            if r.short_term_limit_million_vnd is not None:
                tot_limit += r.short_term_limit_million_vnd
            if r.short_term_debt_vnd_million is not None:
                tot_vnd += r.short_term_debt_vnd_million
            if r.short_term_debt_usd_million is not None:
                tot_usd += r.short_term_debt_usd_million
            if r.medium_long_term_debt_million is not None:
                tot_tdh += r.medium_long_term_debt_million
            if r.total_debt_million is not None:
                tot_debt += r.total_debt_million

            debt_grp_str = f"Nhóm {r.debt_group.value}" if r.debt_group is not None else "-"
            rows_e.append([
                str(r.stt),
                r.bank_name,
                fmt_curr(r.short_term_limit_million_vnd),
                fmt_curr(r.short_term_debt_vnd_million),
                fmt_curr(r.short_term_debt_usd_million),
                fmt_curr(r.medium_long_term_debt_million),
                fmt_curr(r.total_debt_million),
                r.collateral_description or "-",
                debt_grp_str
            ])

        # Dòng Tổng cộng
        rows_e.append([
            "",
            "TỔNG CỘNG",
            fmt_curr(tot_limit),
            fmt_curr(tot_vnd),
            fmt_curr(tot_usd),
            fmt_curr(tot_tdh),
            fmt_curr(tot_debt),
            "",
            f"Nhóm {data.highest_debt_group.value}"
        ])

        cls._create_styled_table(
            doc, headers_e, rows_e,
            col_widths=[0.4, 2.0, 1.1, 1.0, 1.0, 1.0, 1.1, 1.6, 0.8],
            align_right_cols=[2, 3, 4, 5, 6]
        )

        # 2. Nhận xét đánh giá giao dịch tín dụng của RM
        add_h2("2. Nhận xét giao dịch tín dụng & Lịch sử trả nợ")
        if data.is_overdue_12m is False:
            overdue_str = "Trong 12 tháng gần nhất, Khách hàng không có nợ quá hạn, thanh toán đầy đủ gốc và lãi."
        elif data.is_overdue_12m is True:
            overdue_str = f"Lịch sử quá hạn: {data.overdue_explanation}" if data.overdue_explanation else "Có phát sinh nợ quá hạn trong 12 tháng gần nhất (cần RM bổ sung giải trình)."
        else:
            overdue_str = "Chưa ghi nhận thông tin xác nhận về lịch sử nợ quá hạn 12 tháng trên CIC (cần RM rà soát bổ sung)."
        add_body(overdue_str, bold_prefix="- Lịch sử thanh toán nợ 12 tháng: ")

        if data.rm_credit_assessment:
            add_body(data.rm_credit_assessment, bold_prefix="- Đánh giá của ĐVKD: ")

        add_body(f"Dư nợ cho vay hiện hữu: {fmt_curr(data.loan_outstanding_at_msb_million)} triệu đồng. Tổng dư nợ tín dụng tại MSB: {fmt_curr(data.total_credit_exposure_at_msb_million)} triệu đồng.", bold_prefix="- Quan hệ tín dụng tại MSB: ")

        # 3. Giao dịch phái sinh (Generic data-gap notice if absent)
        add_h2("3. Thông tin giao dịch phái sinh tại MSB và các TCTD")
        deriv_str = data.derivative_transactions_info if data.derivative_transactions_info else "Chưa ghi nhận thông tin giao dịch phái sinh trên Báo cáo CIC (cần RM rà soát bổ sung)."
        add_body(deriv_str, bold_prefix="- Trạng thái giao dịch: ")

    @classmethod
    def generate_docx(cls, data: SectionEData, output_path: str) -> str:
        """Tạo file Word độc lập cho Phần E."""
        doc = docx.Document()
        for s in doc.sections:
            s.top_margin = Inches(0.7)
            s.bottom_margin = Inches(0.7)
            s.left_margin = Inches(0.75)
            s.right_margin = Inches(0.75)

        p_title = doc.add_paragraph()
        r_title = p_title.add_run("PHẦN E. THÔNG TIN QUAN HỆ TÍN DỤNG CỦA KHÁCH HÀNG")
        r_title.bold = True
        r_title.font.name = "Times New Roman"
        r_title.font.size = Pt(12)
        r_title.font.color.rgb = RGBColor(0, 32, 96)

        cls.render_to_document(doc, data)

        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        doc.save(output_path)
        return output_path
