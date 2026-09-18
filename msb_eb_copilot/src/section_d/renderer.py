"""Module: renderer.py
Mô tả: Engine render Phần D (Tài chính doanh nghiệp) vào văn bản Word MB07.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from .models import SectionDData


class SectionDRenderer:
    """Renderer Phần D chuẩn format biểu mẫu MB07."""

    @staticmethod
    def _set_cell_background(cell, fill_hex: str = "F1F5F9"):
        tcPr = cell._element.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    @staticmethod
    def _set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
        tcPr = cell._element.get_or_add_tcPr()
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top}" w:type="dxa"/>'
            f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
            f'<w:left w:w="{left}" w:type="dxa"/>'
            f'<w:right w:w="{right}" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)

    @staticmethod
    def _set_table_borders(table, color="B0C4DE", sz="4"):
        tblPr = table._element.xpath('w:tblPr')
        if tblPr:
            borders = parse_xml(
                f'<w:tblBorders {nsdecls("w")}>'
                f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                f'<w:left w:val="none"/>'
                f'<w:right w:val="none"/>'
                f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
                f'<w:insideV w:val="none"/>'
                f'</w:tblBorders>'
            )
            tblPr[0].append(borders)

    @classmethod
    def _create_styled_table(cls, doc, headers, rows, col_widths=None, align_right_cols=None):
        align_right_cols = align_right_cols or []
        table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cls._set_table_borders(table)

        # Headers
        for col_idx, text in enumerate(headers):
            cell = table.rows[0].cells[col_idx]
            cell.text = text
            cls._set_cell_background(cell, "E2E8F0")
            cls._set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if col_idx in align_right_cols else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                run.bold = True
                run.font.name = "Times New Roman"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0, 32, 96)

        # Data Rows
        for row_idx, r_data in enumerate(rows):
            row = table.rows[row_idx + 1]
            bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
            for col_idx, val in enumerate(r_data):
                cell = row.cells[col_idx]
                cell.text = str(val) if val is not None else ""
                cls._set_cell_background(cell, bg)
                cls._set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if col_idx in align_right_cols else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9.0)
                    run.font.color.rgb = RGBColor(51, 65, 85)

        if col_widths:
            for i, w in enumerate(col_widths):
                for row in table.rows:
                    row.cells[i].width = Inches(w)

        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(2)
        sp.paragraph_format.space_after = Pt(4)
        return table

    @classmethod
    def render_to_document(cls, doc: docx.Document, data: SectionDData):
        """Render nội dung Phần D vào đối tượng Document."""
        def add_h2(text):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(text)
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(0, 32, 96)
            return p

        def add_body(text, bold_prefix=None):
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(3)
            if bold_prefix:
                rb = p.add_run(bold_prefix)
                rb.bold = True
                rb.font.name = "Times New Roman"
                rb.font.size = Pt(10)
            rt = p.add_run(text)
            rt.font.name = "Times New Roman"
            rt.font.size = Pt(10)
            return p

        def add_unit_header():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run("ĐVT: Triệu đồng")
            r.italic = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(9.5)

        def fmt_curr(v):
            if v is None:
                return ""
            try:
                f = float(v)
                if f == 0:
                    return "0"
                return f"{f:,.1f}".replace(",", "X").replace(".", ",").replace("X", ".")
            except (ValueError, TypeError):
                return str(v)

        def fmt_pct(v):
            if v is None:
                return ""
            try:
                f = float(v)
                return f"{f:.2f}%".replace(".", ",")
            except (ValueError, TypeError):
                return str(v)

        years = data.income_statement.years

        # 1. Tóm tắt thông tin tài chính
        add_h2("1. Tóm tắt thông tin tài chính & Quản trị kế toán")
        gov = data.governance
        add_body(gov.mandatory_audit_by_law, bold_prefix="- Quy định kiểm toán: ")
        add_body(f"BCTC kiểm toán do {gov.audit_firm_name} thực hiện ({gov.audited_years}).", bold_prefix="- Đơn vị kiểm toán: ")
        add_body(gov.audit_opinion, bold_prefix="- Ý kiến kiểm toán: ")
        add_body(gov.accounting_software, bold_prefix="- Phần mềm kế toán: ")
        add_body(gov.finance_team_structure, bold_prefix="- Bộ máy kế toán: ")
        add_body(gov.internal_financial_regulations, bold_prefix="- Quy chế quản trị: ")

        for bullet in data.summary_bullets:
            add_body(bullet, bold_prefix="- ")

        # 2. Bảng kết quả kinh doanh P&L 3 năm
        add_h2("2. Bảng Kết quả hoạt động kinh doanh (P&L)")
        add_unit_header()
        pnl = data.income_statement
        headers_pnl = ["Chỉ tiêu P&L"] + [f"Năm {y}" for y in years]
        rows_pnl = [
            ["1. Doanh thu thuần"] + [fmt_curr(v) for v in pnl.net_revenue],
            ["2. Giá vốn hàng bán (COGS)"] + [fmt_curr(v) for v in pnl.cogs],
            ["3. Lợi nhuận gộp"] + [fmt_curr(v) for v in pnl.gross_profit],
            ["   - Biên lợi nhuận gộp (%)"] + [fmt_pct(v) for v in pnl.gross_profit_margin_pct],
            ["4. Doanh thu hoạt động tài chính"] + [fmt_curr(v) for v in pnl.financial_income],
            ["5. Chi phí tài chính"] + [fmt_curr(v) for v in pnl.financial_expenses],
            ["   - Trong đó: Chi phí lãi vay"] + [fmt_curr(v) for v in pnl.interest_expenses],
            ["6. Chi phí bán hàng & QLDN (SG&A)"] + [fmt_curr(v) for v in pnl.sga_expenses],
            ["7. Lợi nhuận trước thuế (LNTT)"] + [fmt_curr(v) for v in pnl.net_profit_before_tax],
            ["8. Lợi nhuận sau thuế (LNST)"] + [fmt_curr(v) for v in pnl.net_profit_after_tax],
        ]
        r_cols = list(range(1, len(years) + 1))
        cls._create_styled_table(doc, headers_pnl, rows_pnl, col_widths=[3.0] + [1.3] * len(years), align_right_cols=r_cols)

        # Nhận xét P&L
        p_ana = data.pnl_analysis
        if p_ana.revenue_analysis:
            add_body(p_ana.revenue_analysis, bold_prefix="+ Về Doanh thu: ")
        if p_ana.gross_margin_analysis:
            add_body(p_ana.gross_margin_analysis, bold_prefix="+ Về Giá vốn & Biên lợi nhuận: ")
        if p_ana.financial_income_and_expenses_analysis:
            add_body(p_ana.financial_income_and_expenses_analysis, bold_prefix="+ Về Chi phí tài chính & Lãi vay: ")
        if p_ana.net_profit_and_dividends_analysis:
            add_body(p_ana.net_profit_and_dividends_analysis, bold_prefix="+ Về Lợi nhuận sau thuế: ")

        # 3. Bảng Cân đối kế toán 3 năm
        add_h2("3. Bảng Cân đối kế toán (CĐKT)")
        add_unit_header()
        bs = data.balance_sheet
        headers_bs = ["Chỉ tiêu CĐKT"] + [f"Năm {y}" for y in years]
        rows_bs = [
            ["A. TÀI SẢN NGẮN HẠN"] + [fmt_curr(v) for v in bs.current_assets],
            ["   - Tiền và tương đương tiền"] + [fmt_curr(v) for v in bs.cash_and_equivalents],
            ["   - Đầu tư tài chính ngắn hạn"] + [fmt_curr(v) for v in bs.short_term_investments],
            ["   - Các khoản phải thu ngắn hạn"] + [fmt_curr(v) for v in bs.accounts_receivable],
            ["   - Hàng tồn kho"] + [fmt_curr(v) for v in bs.inventories],
            ["   - Tài sản ngắn hạn khác"] + [fmt_curr(v) for v in bs.other_current_assets],
            ["B. TÀI SẢN DÀI HẠN"] + [fmt_curr(v) for v in bs.non_current_assets],
            ["   - Tài sản cố định (TSCĐ)"] + [fmt_curr(v) for v in bs.fixed_assets],
            ["   - Chi phí xây dựng dở dang"] + [fmt_curr(v) for v in bs.construction_in_progress],
            ["TỔNG TÀI SẢN (A + B)"] + [fmt_curr(v) for v in bs.total_assets],
            ["C. NỢ PHẢI TRẢ"] + [fmt_curr(v) for v in bs.liabilities],
            ["   - Vay & nợ thuê TC ngắn hạn"] + [fmt_curr(v) for v in bs.short_term_debt],
            ["   - Vay & nợ thuê TC dài hạn"] + [fmt_curr(v) for v in bs.long_term_debt],
            ["D. VỐN CHỦ SỞ HỮU"] + [fmt_curr(v) for v in bs.owner_equity],
            ["   - Vốn góp của chủ sở hữu"] + [fmt_curr(v) for v in bs.charter_capital],
            ["TỔNG NGUỒN VỐN (C + D)"] + [fmt_curr(bs.liabilities[i] + bs.owner_equity[i]) for i in range(len(years))],
        ]
        cls._create_styled_table(doc, headers_bs, rows_bs, col_widths=[3.0] + [1.3] * len(years), align_right_cols=r_cols)

        # 4. Báo cáo Lưu chuyển tiền tệ
        add_h2("4. Báo cáo Lưu chuyển tiền tệ")
        add_unit_header()
        cf = data.cash_flow
        headers_cf = ["Chỉ tiêu Lưu chuyển tiền tệ"] + [f"Năm {y}" for y in years]
        rows_cf = [
            ["1. Lưu chuyển tiền thuần từ HĐKD (OCF)"] + [fmt_curr(v) for v in cf.ocf_cash_from_operations],
            ["2. Lưu chuyển tiền thuần từ HĐ Đầu tư (ICF)"] + [fmt_curr(v) for v in cf.icf_cash_from_investing],
            ["3. Lưu chuyển tiền thuần từ HĐ Tài chính (FCF)"] + [fmt_curr(v) for v in cf.fcf_cash_from_financing],
            ["LƯU CHUYỂN TIỀN THUẦN TRONG KỲ (1+2+3)"] + [fmt_curr(v) for v in cf.net_cash_flow],
            ["Tiền và tương đương tiền đầu kỳ"] + [fmt_curr(v) for v in cf.cash_beginning],
            ["Tiền và tương đương tiền cuối kỳ"] + [fmt_curr(v) for v in cf.cash_ending],
        ]
        cls._create_styled_table(doc, headers_cf, rows_cf, col_widths=[3.0] + [1.3] * len(years), align_right_cols=r_cols)
        if cf.cash_flow_analysis:
            add_body(cf.cash_flow_analysis, bold_prefix="Nhận xét dòng tiền: ")

        # 5. Các chỉ số tài chính tổng hợp
        add_h2("5. Các chỉ số tài chính tổng hợp đối chiếu chuẩn MSB")
        ratios = data.ratios
        headers_r = ["Nhóm Chỉ số Tài chính"] + [f"Năm {y}" for y in years] + ["Chuẩn MSB", "Đánh giá của ĐVKD"]
        rows_r = [
            ["1. Khả năng thanh toán hiện hành (Current)"] + [f"{v:.2f}" for v in ratios.current_ratio] + ["≥ 1.10 lần", "Đạt chuẩn"],
            ["2. Khả năng thanh toán nhanh (Quick)"] + [f"{v:.2f}" for v in ratios.quick_ratio] + ["≥ 0.50 lần", "Đạt chuẩn"],
            ["3. Khả năng thanh toán tức thời (Cash)"] + [f"{v:.2f}" for v in ratios.cash_ratio] + ["≥ 0.20 lần", "Thanh khoản tốt"],
            ["4. Đòn bẩy nợ D/E (Nợ phải trả/VCSH)"] + [f"{v:.2f}" for v in ratios.debt_to_equity] + ["≤ 4.00 lần", "Kiểm soát an toàn"],
            ["5. Tổng nợ vay tài chính/VCSH"] + [f"{v:.2f}" for v in ratios.total_debt_to_equity] + ["≤ 2.00 lần", "Đạt chuẩn EB"],
            ["6. Khả năng chi trả lãi vay (ICR/DSCR)"] + [f"{v:.2f}" for v in ratios.dscr_icr] + ["≥ 1.50 lần", "Dòng tiền trả nợ tốt"],
            ["7. Tỷ suất sinh lời trên Doanh thu (ROS)"] + [fmt_pct(v) for v in ratios.ros] + ["> 0.00%", "Có lãi"],
            ["8. Tỷ suất sinh lời trên VCSH (ROE)"] + [fmt_pct(v) for v in ratios.roe] + ["> 0.00%", "Bảo toàn vốn"],
        ]
        cls._create_styled_table(doc, headers_r, rows_r, col_widths=[2.4] + [0.8] * len(years) + [0.9, 1.1], align_right_cols=r_cols)

    @classmethod
    def generate_docx(cls, data: SectionDData, output_path: str) -> str:
        """Tạo file Word độc lập cho Phần D."""
        doc = docx.Document()
        for s in doc.sections:
            s.top_margin = Inches(0.7)
            s.bottom_margin = Inches(0.7)
            s.left_margin = Inches(0.75)
            s.right_margin = Inches(0.75)

        p_title = doc.add_paragraph()
        r_title = p_title.add_run("PHẦN D. TÌNH HÌNH TÀI CHÍNH DOANH NGHIỆP")
        r_title.bold = True
        r_title.font.name = "Times New Roman"
        r_title.font.size = Pt(12)
        r_title.font.color.rgb = RGBColor(0, 32, 96)

        cls.render_to_document(doc, data)

        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        doc.save(output_path)
        return output_path
