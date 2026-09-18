"""Unit tests cho Phân hệ C - Hoạt động kinh doanh & Chuỗi cung ứng."""

import unittest
import os
import docx

from msb_eb_copilot.src.section_c import (
    BusinessModelType,
    BlacklistStatus,
    ConcentrationRiskLevel,
    ShareholderInfo,
    ManagementMember,
    ProductInfo,
    WarehouseInfo,
    EquipmentInfo,
    SupplierInfo,
    CustomerInfo,
    SectionCData,
    SectionCValidator,
    SectionCRenderer,
)


class TestSectionCValidator(unittest.TestCase):
    """Kiểm thử tính toàn vẹn và các quy tắc kiểm định của Section C."""

    def setUp(self):
        self.sample_data = SectionCData(
            customer_name="CÔNG TY TNHH THÉP TÂY ĐÔ",
            history_narrative="Thành lập năm 1995, hơn 30 năm kinh nghiệm trong ngành luyện cán thép.",
            parent_company_or_owner="Các thành viên góp vốn tư nhân",
            major_shareholders=[
                ShareholderInfo(stt=1, shareholder_name="Huỳnh Trung Quang", id_tax_code="024088001234", ownership_percentage=55.0, contributed_capital_million_vnd=275000.0),
                ShareholderInfo(stt=2, shareholder_name="Nguyễn Thị Kim Loan", id_tax_code="024190005678", ownership_percentage=45.0, contributed_capital_million_vnd=225000.0),
            ],
            blacklist_status=BlacklistStatus.KHONG_VI_PHAM,
            management_members=[
                ManagementMember(position="Tổng Giám đốc", full_name="Huỳnh Trung Quang", profile_summary="30 năm kinh nghiệm điều hành", years_at_company=15),
                ManagementMember(position="Kế toán trưởng", full_name="Trần Văn Nam", profile_summary="Cử nhân Kế toán, 12 năm kinh nghiệm", years_at_company=8),
            ],
            business_model=BusinessModelType.SAN_XUAT,
            products=[
                ProductInfo(stt=1, product_name="Phôi thép", brand_name="Thép Tây Đô", revenue_share_percentage=65.0),
                ProductInfo(stt=2, product_name="Thép xây dựng", brand_name="Thép Tây Đô", revenue_share_percentage=35.0),
            ],
            suppliers=[
                SupplierInfo(stt=1, supplier_name="CTCP Thép Pomina", supplied_goods="Thép phế liệu", purchase_share_percentage=25.0, payment_terms="L/C 90 ngày"),
                SupplierInfo(stt=2, supplier_name="Mitsui & Co", supplied_goods="Thép phế nhập khẩu", purchase_share_percentage=15.0, payment_terms="L/C at sight"),
            ],
            customers=[
                CustomerInfo(stt=1, customer_name="Công ty CP Thép Chín Rồng", product_purchased="Phôi thép", revenue_share_percentage=20.0, credit_terms="Trả góp 15 ngày"),
                CustomerInfo(stt=2, customer_name="Chip Mong Group", product_purchased="Thép thanh", revenue_share_percentage=15.0, credit_terms="L/C UPAS"),
            ],
            rm_supply_chain_assessment="Chuỗi cung ứng ổn định, nguồn cung phế liệu đa dạng.",
            rm_credit_risk_mitigation="Kiểm soát giải ngân theo hoá đơn mua hàng."
        )

    def test_valid_section_c_passes(self):
        """Hồ sơ chuẩn mực phải vượt qua kiểm định không có lỗi."""
        res = SectionCValidator.validate(self.sample_data)
        self.assertTrue(res.is_valid)
        self.assertEqual(len(res.errors), 0)
        self.assertEqual(len(res.concentration_risks), 0)

    def test_blacklist_blocking_error(self):
        """Khách hàng thuộc Blacklist phải bị chặn lập tức."""
        self.sample_data.blacklist_status = BlacklistStatus.THUOC_BLACKLIST
        res = SectionCValidator.validate(self.sample_data)
        self.assertFalse(res.is_valid)
        self.assertTrue(any("BLACKLIST" in e for e in res.errors))

    def test_excess_shareholder_ownership_error(self):
        """Tổng tỷ lệ sở hữu cổ đông > 100% phải báo lỗi."""
        self.sample_data.major_shareholders.append(
            ShareholderInfo(stt=3, shareholder_name="Cổ đông C", id_tax_code="001", ownership_percentage=10.0, contributed_capital_million_vnd=50000.0)
        )
        res = SectionCValidator.validate(self.sample_data)
        self.assertFalse(res.is_valid)
        self.assertTrue(any("vượt quá 100%" in e for e in res.errors))

    def test_supplier_concentration_warning(self):
        """NCC chiếm >= 40% chi phí mua phải kích hoạt cảnh báo rủi ro tập trung."""
        self.sample_data.suppliers[0].purchase_share_percentage = 45.0
        res = SectionCValidator.validate(self.sample_data)
        self.assertTrue(res.is_valid)  # Vẫn valid nhưng có cảnh báo
        self.assertIn(ConcentrationRiskLevel.CANH_BAO_TAP_TRUNG_NCC, res.concentration_risks)
        self.assertTrue(any("RỦI RO TẬP TRUNG ĐẦU VÀO" in w for w in res.warnings))

    def test_customer_concentration_warning(self):
        """Khách hàng chiếm >= 30% doanh thu phải kích hoạt cảnh báo rủi ro tập trung."""
        self.sample_data.customers[0].revenue_share_percentage = 35.0
        res = SectionCValidator.validate(self.sample_data)
        self.assertTrue(res.is_valid)
        self.assertIn(ConcentrationRiskLevel.CANH_BAO_TAP_TRUNG_KH, res.concentration_risks)
        self.assertTrue(any("RỦI RO TẬP TRUNG ĐẦU RA" in w for w in res.warnings))

    def test_renderer_generates_docx(self):
        """Kiểm tra renderer tạo file Word Phần C hợp lệ."""
        test_out = os.path.join("output", "test_section_c_out.docx")
        os.makedirs("output", exist_ok=True)
        SectionCRenderer.generate_docx(self.sample_data, test_out)
        self.assertTrue(os.path.exists(test_out))
        
        doc = docx.Document(test_out)
        self.assertGreater(len(doc.paragraphs), 5)
        self.assertGreater(len(doc.tables), 3)


if __name__ == "__main__":
    unittest.main()
