"""Unit tests cho Phân hệ E - Quan hệ tín dụng & Báo cáo CIC đa ngân hàng."""

import unittest
import os
from decimal import Decimal
import docx

from msb_eb_copilot.src.section_e import (
    DebtGroup,
    CollateralType,
    CreditInstitutionRelation,
    SectionEData,
    SectionEValidator,
    SectionERenderer,
    link_section_e_to_session_a,
    get_other_debt_for_section_d,
)
from msb_eb_copilot.src.section_a.review_session import SectionAReviewSession


class TestSectionECIC(unittest.TestCase):
    """Kiểm thử tính toàn vẹn và các quy tắc kiểm định quan hệ tín dụng Phần E."""

    def setUp(self):
        self.sample_data = SectionEData(
            customer_name="CÔNG TY TNHH THÉP TÂY ĐÔ",
            cic_report_date="15/08/2025",
            relations=[
                CreditInstitutionRelation(
                    stt=1,
                    bank_name="BIDV - CN Cần Thơ",
                    short_term_limit_million_vnd=400000.0,
                    short_term_debt_vnd_million=250000.0,
                    short_term_debt_usd_million=0.0,
                    medium_long_term_debt_million=100000.0,
                    total_debt_million=350000.0,
                    collateral_description="BĐS nhà xưởng",
                    debt_group=DebtGroup.NHOM_1_DU_TIEU_CHUAN
                ),
                CreditInstitutionRelation(
                    stt=2,
                    bank_name="VietinBank - CN Tây Đô",
                    short_term_limit_million_vnd=300000.0,
                    short_term_debt_vnd_million=180000.0,
                    short_term_debt_usd_million=0.0,
                    medium_long_term_debt_million=0.0,
                    total_debt_million=180000.0,
                    collateral_description="HĐTG, Phôi thép",
                    debt_group=DebtGroup.NHOM_1_DU_TIEU_CHUAN
                ),
                CreditInstitutionRelation(
                    stt=3,
                    bank_name="MSB - CN Cần Thơ",
                    short_term_limit_million_vnd=6500.0,
                    short_term_debt_vnd_million=6500.0,
                    short_term_debt_usd_million=0.0,
                    medium_long_term_debt_million=0.0,
                    total_debt_million=6500.0,
                    collateral_description="Bảo lãnh Quỹ Advance",
                    debt_group=DebtGroup.NHOM_1_DU_TIEU_CHUAN
                ),
            ],
            loan_outstanding_at_msb_million=6500.0,
            total_credit_exposure_at_msb_million=6500.0,
            is_overdue_12m=False,
            rm_credit_assessment="Khách hàng trả nợ gốc và lãi đúng hạn, không có nợ quá hạn."
        )

    def test_valid_section_e_passes(self):
        """Hồ sơ tín dụng không nợ xấu phải vượt qua thẩm định."""
        res = SectionEValidator.validate(self.sample_data)
        self.assertTrue(res.is_valid)
        self.assertEqual(len(res.errors), 0)
        self.assertFalse(res.has_bad_debt)
        self.assertFalse(res.has_special_mention_debt)

    def test_bad_debt_group_3_blocking_error(self):
        """Khách hàng có nợ Nhóm 3 (Nợ xấu) phải bị chặn cấp tín dụng."""
        self.sample_data.relations[0].debt_group = DebtGroup.NHOM_3_DUOI_TIEU_CHUAN
        res = SectionEValidator.validate(self.sample_data)
        self.assertFalse(res.is_valid)
        self.assertTrue(res.has_bad_debt)
        self.assertTrue(any("NỢ XẤU" in e for e in res.errors))

    def test_special_mention_debt_group_2_warning(self):
        """Khách hàng có nợ Nhóm 2 phải kích hoạt cảnh báo yêu cầu RM giải trình."""
        self.sample_data.relations[1].debt_group = DebtGroup.NHOM_2_CAN_CHU_Y
        res = SectionEValidator.validate(self.sample_data)
        self.assertTrue(res.is_valid)  # Vẫn xét duyệt có điều kiện
        self.assertTrue(res.has_special_mention_debt)
        self.assertTrue(any("CẢNH BÁO NỢ CẦN CHÚ Ý" in w or "Nhóm 2" in w for w in res.warnings))

    def test_unexplained_overdue_error(self):
        """Có phát sinh nợ quá hạn trong 12 tháng mà không có giải trình phải báo lỗi."""
        self.sample_data.is_overdue_12m = True
        self.sample_data.overdue_explanation = ""
        res = SectionEValidator.validate(self.sample_data)
        self.assertFalse(res.is_valid)
        self.assertTrue(any("quá hạn" in e for e in res.errors))

    def test_cross_section_linking_to_section_a(self):
        """Kiểm tra liên kết chéo tự động từ Section E sang Section A."""
        session_a = SectionAReviewSession("CASE-TEST-CROSS-LINK")
        link_section_e_to_session_a(self.sample_data, session_a)

        fact_loan = session_a.facts["credit_relation.loan_outstanding_at_msb"]
        fact_exp = session_a.facts["credit_relation.total_credit_exposure_at_msb"]

        self.assertEqual(fact_loan.value.value, Decimal("6500.00"))
        self.assertEqual(fact_exp.value.value, Decimal("6500.00"))

    def test_get_other_debt_for_section_d(self):
        """Kiểm tra tính tổng nợ các ngân hàng khác loại trừ MSB để cấp cho MB09."""
        other_debt_vnd = get_other_debt_for_section_d(self.sample_data)
        # BIDV: 350.000 trđ + VietinBank: 180.000 trđ = 530.000 trđ = 530 tỷ VND
        expected_vnd = 530000.0 * 1_000_000.0
        self.assertEqual(other_debt_vnd, expected_vnd)

    def test_renderer_generates_docx(self):
        """Kiểm tra renderer tạo file Word Bảng 07 Phần E hợp lệ."""
        test_out = os.path.join("output", "test_section_e_out.docx")
        os.makedirs("output", exist_ok=True)
        SectionERenderer.generate_docx(self.sample_data, test_out)
        self.assertTrue(os.path.exists(test_out))
        
        doc = docx.Document(test_out)
        self.assertGreater(len(doc.paragraphs), 5)
        self.assertGreater(len(doc.tables), 0)


if __name__ == "__main__":
    unittest.main()
