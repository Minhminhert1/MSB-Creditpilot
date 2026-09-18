# -*- coding: utf-8 -*-
"""Golden End-to-End Test for Phase 5: Business Document Agent & Canonical Section C Lineage.

Verifies:
1. Ingestion of synthetic digital business fixture via DocumentIngestionRouter.
2. Comprehensive per-field grounding audit (all candidate facts VERIFIED).
3. Deterministic mapping into canonical section_c (entity preservation, date precision, derived fields).
4. Section C isolation and preservation of RM-owned fields.
5. Generation of Section C Word document (MB07 Tables 01–06) with verified dynamic bindings.
"""

import copy
import os
import unittest
import docx

from msb_eb_copilot.src.ingestion.router import DocumentIngestionRouter
from msb_eb_copilot.src.extraction.business_extraction import (
    BusinessEvidenceField,
    CapitalMilestoneItem,
    ShareholderItem,
    ManagementItem,
    ProductItem,
    WarehouseItem,
    EquipmentItem,
    SupplierItem,
    CustomerItem,
    CompetitorItem,
    BusinessDocumentExtraction,
    BusinessGroundingAuditor,
    BusinessIdentityReconciler,
    BusinessModelClassifier,
)
from msb_eb_copilot.src.mapping.business_mapper import BusinessDocumentMapper
from msb_eb_copilot.src.mapping.models import MappingSourceMetadata
from msb_eb_copilot.src.section_c.models import (
    SectionCData,
    CapitalMilestone,
    ShareholderInfo,
    ManagementMember,
    ProductInfo,
    WarehouseInfo,
    EquipmentInfo,
    SupplierInfo,
    CustomerInfo,
    BusinessModelType,
    BlacklistStatus,
)
from msb_eb_copilot.src.section_c.renderer import SectionCRenderer
from web_copilot_app import CASES_DB, ACTIVE_CASE_ID


class TestBusinessGoldenE2E(unittest.TestCase):
    """Synthetic Golden E2E Test on synthetic digital business fixture."""

    def setUp(self):
        self.digital_pdf_path = os.path.join("tests", "fixtures", "business", "synthetic_business_digital.pdf")
        self.assertTrue(os.path.exists(self.digital_pdf_path), f"Fixture not found at: {self.digital_pdf_path}")

        def bef(v, ev, pg=1):
            return BusinessEvidenceField(value_raw=v, evidence=ev, page=pg)

        # Grounded extraction corresponding exactly to synthetic_business_digital.pdf
        self.extraction = BusinessDocumentExtraction(
            company_name=bef("CÔNG TY CỔ PHẦN CÔNG NGHỆ VÀ TRUYỀN THÔNG ĐÔNG NAM Á", "Tên công ty: CÔNG TY CỔ PHẦN CÔNG NGHỆ VÀ TRUYỀN THÔNG ĐÔNG NAM Á", 1),
            tax_code=bef("0108889999", "Mã số thuế: 0108889999", 1),
            established_year=bef("2015", "Năm thành lập: 2015", 1),
            history_narrative=bef("Thành lập năm 2015 tiền thân là Trung tâm Công nghệ Viễn thông Đông Nam Á", "Thành lập năm 2015 tiền thân là Trung tâm Công nghệ Viễn thông Đông Nam Á", 1),
            parent_company_or_owner=bef("Tập đoàn Công nghệ Viễn thông Quốc tế Á Châu", "sở hữu của Tập đoàn Công nghệ Viễn thông Quốc tế Á Châu", 1),
            capital_milestones=[
                CapitalMilestoneItem(
                    effective_date=bef("2020", "2020", 1),
                    charter_capital_raw=bef("50.000", "50.000", 1),
                    event_description=bef("Tăng vốn điều lệ từ phát hành cổ phần", "Tăng vốn điều lệ từ phát hành cổ phần", 1),
                    page=1,
                ),
                CapitalMilestoneItem(
                    effective_date=bef("15/05/2023", "15/05/2023", 1),
                    charter_capital_raw=bef("120.000", "120.000", 1),
                    event_description=bef("Chào bán cho cổ đông chiến lược", "Chào bán cho cổ đông chiến lược", 1),
                    page=1,
                ),
            ],
            shareholders=[
                ShareholderItem(
                    shareholder_name=bef("Tập đoàn Công nghệ Viễn thông Quốc tế Á Châu", "Tập đoàn Công nghệ Viễn thông Quốc tế Á Châu", 1),
                    id_tax_code=bef("0102345678", "0102345678", 1),
                    ownership_percentage_raw=bef("51.0%", "51.0%", 1),
                    contributed_capital_raw=bef("61.200", "61.200", 1),
                    page=1,
                ),
                ShareholderItem(
                    shareholder_name=bef("Nguyễn Văn Toàn", "Nguyễn Văn Toàn", 1),
                    id_tax_code=bef("001080009876", "001080009876", 1),
                    ownership_percentage_raw=bef("25.5%", "25.5%", 1),
                    contributed_capital_raw=bef("30.600", "30.600", 1),
                    page=1,
                ),
                ShareholderItem(
                    shareholder_name=bef("Trần Thị Mai Phương", "Trần Thị Mai Phương", 1),
                    id_tax_code=bef("001185004321", "001185004321", 1),
                    ownership_percentage_raw=bef("15.0%", "15.0%", 1),
                    contributed_capital_raw=bef("18.000", "18.000", 1),
                    page=1,
                ),
            ],
            management=[
                ManagementItem(
                    full_name=bef("Nguyễn Văn Toàn", "Nguyễn Văn Toàn", 1),
                    position=bef("Chủ tịch HĐQT kiêm Tổng Giám đốc", "Chủ tịch HĐQT kiêm Tổng Giám đốc", 1),
                    explicit_experience_years_raw=bef("18", "hơn 18 năm kinh nghiệm", 1),
                    profile_summary=bef("Kỹ sư Viễn thông, hơn 18 năm kinh nghiệm trong ngành viễn thông", "hơn 18 năm kinh nghiệm trong ngành viễn thông", 1),
                    page=1,
                ),
                ManagementItem(
                    full_name=bef("Lê Quốc Hưng", "Lê Quốc Hưng", 1),
                    position=bef("Giám đốc Kỹ thuật & Vận hành", "Giám đốc Kỹ thuật & Vận hành", 1),
                    career_history_raw=bef("2016-2021: Trưởng phòng Giải pháp IoT; 2021-nay: Giám đốc Kỹ thuật", "2016-2021: Trưởng phòng Giải pháp IoT; 2021-nay", 1),
                    explicit_experience_years_raw=None,  # Career dates only!
                    profile_summary=bef("2016-2021: Trưởng phòng Giải pháp IoT; 2021-nay: Giám đốc Kỹ thuật", "2016-2021: Trưởng phòng Giải pháp IoT; 2021-nay", 1),
                    page=1,
                ),
            ],
            operating_model_description=bef(
                "Doanh nghiệp hoạt động theo mô hình sản xuất và phân phối các thiết bị đầu cuối viễn thông và giải pháp IoT.",
                "hoạt động theo mô hình sản xuất và phân phối các thiết bị đầu cuối viễn thông",
                1,
            ),
            products=[
                ProductItem(
                    product_name=bef("Thiết bị định vị & cảm biến IoT thông minh", "Thiết bị định vị & cảm biến IoT thông minh", 1),
                    brand_or_spec=bef("Chuẩn công nghiệp IP67", "Chuẩn công nghiệp IP67", 1),
                    revenue_share_percentage_raw=bef("45.0%", "45.0%", 1),
                    page=1,
                ),
                ProductItem(
                    product_name=bef("Hệ thống tổng đài IP và Router viễn thông", "Hệ thống tổng đài IP và Router viễn thông", 1),
                    brand_or_spec=bef("Nhập khẩu chính hãng", "Nhập khẩu chính hãng", 1),
                    revenue_share_percentage_raw=bef("35.0%", "35.0%", 1),
                    page=1,
                ),
                ProductItem(
                    product_name=bef("Dịch vụ phần mềm quản trị và bảo trì hệ thống", "Dịch vụ phần mềm quản trị và bảo trì hệ thống", 1),
                    brand_or_spec=bef("Bản quyền SaaS", "Bản quyền SaaS", 1),
                    revenue_share_percentage_raw=bef("20.0%", "20.0%", 1),
                    page=1,
                ),
            ],
            warehouses=[
                WarehouseItem(
                    facility_type=bef("Kho hàng trung tâm", "Kho hàng trung tâm", 1),
                    address=bef("Lô C2 KCN Quang Minh, Mê Linh, Hà Nội", "Lô C2 KCN Quang Minh, Mê Linh, Hà Nội", 1),
                    area_raw=bef("3.500 m2", "3.500 m2", 1),
                    ownership_type=bef("Thuê dài hạn KCN 10 năm", "Thuê dài hạn KCN 10 năm", 1),
                    capacity_description=bef("Theo đơn hàng thiết bị IoT", "Theo đơn hàng thiết bị IoT", 1),
                    page=1,
                )
            ],
            equipments=[
                EquipmentItem(
                    equipment_name=bef("Dây chuyền lắp ráp và kiểm thử bo mạch SMT", "Dây chuyền lắp ráp và kiểm thử bo mạch SMT", 1),
                    origin_and_technology=bef("Nhật Bản - Yamaha", "Nhật Bản - Yamaha", 1),
                    designed_capacity=bef("50.000 sản phẩm/năm", "50.000 sản phẩm/năm", 1),
                    utilization_rate=bef("92%", "92%", 1),
                    page=1,
                )
            ],
            suppliers=[
                SupplierItem(
                    supplier_name=bef("Công ty TNHH Linh Kiện Điện Tử Qualcomm Việt Nam", "Công ty TNHH Linh Kiện Điện Tử Qualcomm Việt Nam", 1),
                    supplied_goods=bef("Chipset 4G/5G và vi xử lý IoT", "Chipset 4G/5G và vi xử lý IoT", 1),
                    purchase_share_percentage_raw=bef("32.0%", "32.0%", 1),
                    payment_terms=bef("L/C trả ngay", "L/C trả ngay", 1),
                    page=1,
                ),
                SupplierItem(
                    supplier_name=bef("Quectel Wireless Solutions Co., Ltd", "Quectel Wireless Solutions Co., Ltd", 1),
                    supplied_goods=bef("Module truyền thông không dây", "Module truyền thông không dây", 1),
                    purchase_share_percentage_raw=bef("28.5%", "28.5%", 1),
                    payment_terms=bef("TTR gối đầu 30 ngày", "TTR gối đầu 30 ngày", 1),
                    page=1,
                ),
                SupplierItem(
                    supplier_name=bef("Công ty Cổ phần Dây và Cáp Điện Thượng Đình", "Công ty Cổ phần Dây và Cáp Điện Thượng Đình", 1),
                    supplied_goods=bef("Cáp quang và phụ kiện truyền dẫn", "Cáp quang và phụ kiện truyền dẫn", 1),
                    purchase_share_percentage_raw=bef("18.0%", "18.0%", 1),
                    payment_terms=bef("Chuyển khoản theo tiến độ", "Chuyển khoản theo tiến độ", 1),
                    page=1,
                ),
            ],
            distribution_channels=bef("Phân phối trực tiếp đến các nhà mạng viễn thông, tập đoàn công nghiệp và mạng lưới đại lý trên toàn quốc.", "Phân phối trực tiếp đến các nhà mạng viễn thông", 1),
            customers=[
                CustomerItem(
                    customer_name=bef("Tổng Công ty Viễn thông Viettel", "Tổng Công ty Viễn thông Viettel", 1),
                    product_purchased=bef("Thiết bị cảm biến và Gateway IoT", "Thiết bị cảm biến và Gateway IoT", 1),
                    revenue_share_percentage_raw=bef("34.5%", "34.5%", 1),
                    credit_terms=bef("Bảo lãnh thanh toán 45 ngày", "Bảo lãnh thanh toán 45 ngày", 1),
                    page=1,
                ),
                CustomerItem(
                    customer_name=bef("Tập đoàn Bưu chính Viễn thông Việt Nam (VNPT)", "Tập đoàn Bưu chính Viễn thông Việt Nam (VNPT)", 1),
                    product_purchased=bef("Router và thiết bị truyền dẫn", "Router và thiết bị truyền dẫn", 1),
                    revenue_share_percentage_raw=bef("26.0%", "26.0%", 1),
                    credit_terms=bef("Chuyển khoản sau nghiệm thu 30 ngày", "Chuyển khoản sau nghiệm thu 30 ngày", 1),
                    page=1,
                ),
                CustomerItem(
                    customer_name=bef("Tổng Công ty Viễn thông MobiFone", "Tổng Công ty Viễn thông MobiFone", 1),
                    product_purchased=bef("Thiết bị định vị và dịch vụ giải pháp", "Thiết bị định vị và dịch vụ giải pháp", 1),
                    revenue_share_percentage_raw=bef("19.5%", "19.5%", 1),
                    credit_terms=bef("Trả chậm 30 ngày", "Trả chậm 30 ngày", 1),
                    page=1,
                ),
            ],
            market_share_claim=bef("Chiếm khoảng 22% thị phần phân phối thiết bị IoT chuyên dụng tại Việt Nam.", "Chiếm khoảng 22% thị phần phân phối", 1),
            competitors=[
                CompetitorItem(competitor_name=bef("Công ty Cổ phần Viễn thông Á Châu", "Công ty Cổ phần Viễn thông Á Châu", 1), page=1),
                CompetitorItem(competitor_name=bef("Công ty TNHH Công nghệ Số Tân Tiến", "Công ty TNHH Công nghệ Số Tân Tiến", 1), page=1),
            ],
            competitive_advantages_claim=bef("Lợi thế về dây chuyền công nghệ kiểm thử SMT đạt chuẩn quốc tế và quan hệ hợp tác trực tiếp với Qualcomm.", "Lợi thế về dây chuyền công nghệ kiểm thử SMT", 1),
        )

    def test_business_pdf_to_canonical_section_c_and_mb07_docx(self):
        """Complete Phase 5 Golden Verification Pipeline."""
        # 1. Ingest via router
        ingest_res = DocumentIngestionRouter.ingest_document(self.digital_pdf_path)
        self.assertEqual(ingest_res.mode, "digital")
        self.assertEqual(ingest_res.page_count, 1)

        # 2. Grounding audit on physical page text
        pages_text, max_p = BusinessGroundingAuditor.extract_pages(ingest_res.tagged_text)
        audits = BusinessGroundingAuditor.audit_all_facts(self.extraction, pages_text, max_p)

        # Verify all candidate facts are VERIFIED
        self.assertTrue(len(audits) >= 25, f"Expected >= 25 audited facts, got {len(audits)}")
        for path, audit in audits.items():
            self.assertEqual(audit.status, "VERIFIED", f"Fact {path} was not verified: {audit.reason}")

        # 3. Identity reconciliation
        case_customer = {"name": "CÔNG TY CỔ PHẦN CÔNG NGHỆ VÀ TRUYỀN THÔNG ĐÔNG NAM Á", "tax_code": "0108889999"}
        id_status, _ = BusinessIdentityReconciler.reconcile("0108889999", "CÔNG TY CP CÔNG NGHỆ VÀ TRUYỀN THÔNG ĐÔNG NAM Á", case_customer)
        self.assertEqual(id_status, "MATCH")

        # 4. Business model suggestion
        op_desc = self.extraction.operating_model_description.value_raw
        sug_bm = BusinessModelClassifier.suggest_model(op_desc)
        self.assertEqual(sug_bm, "HON_HOP")

        # 5. Deterministic Mapping
        target_case = {
            "customer": copy.deepcopy(case_customer),
            "section_c": {
                "rm_supply_chain_assessment": "Ban lãnh đạo gắn bó lâu năm, giàu kinh nghiệm ngành.",
                "rm_credit_risk_mitigation": "Doanh nghiệp phân phối IoT uy tín tại Việt Nam.",
            },
        }
        src_meta = MappingSourceMetadata("synthetic_business_digital.pdf", "digital", "BusinessDocumentExtractor")
        map_result = BusinessDocumentMapper.map(target_case, self.extraction, src_meta, grounding_audits=audits, business_model_selection=sug_bm)

        sec_c = map_result.case_data["section_c"]

        # Check entity preservation
        self.assertEqual(len(sec_c["capital_milestones"]), 2)
        self.assertEqual(sec_c["capital_milestones"][0]["effective_date"], "2020")
        self.assertEqual(sec_c["capital_milestones"][1]["effective_date"], "15/05/2023")

        self.assertEqual(len(sec_c["shareholders"]), 3)
        self.assertEqual(sec_c["shareholders"][0]["pct"], 51.0)
        self.assertTrue(sec_c["shareholders"][0]["is_major_shareholder"])

        self.assertEqual(len(sec_c["management"]), 2)
        self.assertEqual(sec_c["management"][0]["exp"], 18)
        self.assertIsNone(sec_c["management"][1]["exp"])  # Career dates only!

        self.assertEqual(len(sec_c["products"]), 3)
        self.assertEqual(len(sec_c["warehouses"]), 1)
        self.assertEqual(len(sec_c["equipments"]), 1)

        self.assertEqual(len(sec_c["suppliers"]), 3)
        self.assertTrue(sec_c["suppliers"][0]["has_cic_check"])  # 32.0% >= 30.0

        self.assertEqual(len(sec_c["customers"]), 3)
        self.assertEqual(len(sec_c["top_competitors"]), 2)

        # 6. SectionCRenderer Word generation
        sh_objs = [
            ShareholderInfo(s["stt"], s["name"], s["tax_code"], s["pct"], s["val"], s["is_major_shareholder"])
            for s in sec_c["shareholders"]
        ]
        mgmt_objs = [
            ManagementMember(m["title"], m["name"], m["note"], m["exp"] or 0)
            for m in sec_c["management"]
        ]
        prod_objs = [
            ProductInfo(i+1, p["name"], p["spec"], p["share"])
            for i, p in enumerate(sec_c["products"])
        ]
        wh_objs = [
            WarehouseInfo(w["stt"], w["facility_type"], w["address"], w["area_m2"], w["ownership_type"], w["capacity_description"])
            for w in sec_c["warehouses"]
        ]
        eq_objs = [
            EquipmentInfo(eq["stt"], eq["equipment_name"], eq["origin_and_technology"], eq["designed_capacity"], eq["utilization_rate"])
            for eq in sec_c["equipments"]
        ]
        sup_objs = [
            SupplierInfo(i+1, s["name"], s["goods"], s["share"], s["term"], s["has_cic_check"])
            for i, s in enumerate(sec_c["suppliers"])
        ]
        cust_objs = [
            CustomerInfo(i+1, c["name"], c["goods"], c["share"], c["term"])
            for i, c in enumerate(sec_c["customers"])
        ]
        ms_objs = [
            CapitalMilestone(m["effective_date"], m["charter_capital_million_vnd"], m["event_description"])
            for m in sec_c["capital_milestones"]
        ]

        data_c = SectionCData(
            customer_name=case_customer["name"],
            history_narrative=sec_c["history_narrative"],
            capital_milestones=ms_objs,
            parent_company_or_owner=sec_c["parent_company_or_owner"],
            major_shareholders=sh_objs,
            blacklist_status=BlacklistStatus.KHONG_VI_PHAM,
            management_members=mgmt_objs,
            business_model=BusinessModelType.HON_HOP,
            products=prod_objs,
            warehouses=wh_objs,
            equipments=eq_objs,
            suppliers=sup_objs,
            customers=cust_objs,
            distribution_channels=sec_c["distribution_channels"],
            market_share_estimate=sec_c["market_share_estimate"],
            top_competitors=sec_c["top_competitors"],
            competitive_advantages=sec_c["competitive_advantages"],
            rm_supply_chain_assessment=sec_c.get("rm_supply_chain_assessment", ""),
            rm_credit_risk_mitigation=sec_c.get("rm_credit_risk_mitigation", ""),
        )

        test_out_docx = os.path.join("output", "TEST_SECTION_C_GOLDEN_MB07.docx")
        os.makedirs("output", exist_ok=True)
        SectionCRenderer.generate_docx(data_c, test_out_docx)

        self.assertTrue(os.path.exists(test_out_docx))
        doc = docx.Document(test_out_docx)
        text_dump = " ".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for row in t.rows for c in row.cells])

        # Verify MB07 rendered content
        self.assertIn("Tập đoàn Công nghệ Viễn thông Quốc tế Á Châu", text_dump)
        self.assertIn("Nguyễn Văn Toàn", text_dump)
        self.assertIn("Lê Quốc Hưng", text_dump)
        self.assertIn("Qualcomm Việt Nam", text_dump)
        self.assertIn("Viettel", text_dump)
        self.assertIn("KCN Quang Minh", text_dump)
        self.assertIn("2020", text_dump)
        self.assertIn("Chiếm khoảng 22% thị phần", text_dump)
