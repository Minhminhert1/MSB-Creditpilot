"""Tests for Section A authoritative DOCX mutator (Phases 12 & 13)."""

import os
import shutil
import tempfile
import unittest
import docx
from docx.oxml.ns import qn

from msb_eb_copilot.src.section_a.docx_mutator import SectionADocxMutator


class SectionADocxMutatorTests(unittest.TestCase):
    """Test suite for Phase 12 & 13 DOCX mutator."""

    def setUp(self) -> None:
        self.repo_root = os.path.abspath(
            os.path.join(os.path.dirname(__file__), "..", "..")
        )
        self.orig_template = os.path.join(
            self.repo_root, "MB07 Tờ trình đề xuất cấp tín dụng (ĐVKD) - Tho.docx"
        )
        self.tmp_dir = tempfile.mkdtemp()
        self.working_copy = os.path.join(self.tmp_dir, "test_mb07.docx")
        shutil.copy2(self.orig_template, self.working_copy)
        self.mutator = SectionADocxMutator(self.working_copy)

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp_dir, ignore_errors=True)

    def test_init_from_doc_and_path(self):
        # From path
        m1 = SectionADocxMutator(self.working_copy)
        self.assertIsNotNone(m1.table)
        self.assertEqual(len(m1.table.rows), 30)

        # From Document
        doc = docx.Document(self.working_copy)
        m2 = SectionADocxMutator(doc)
        self.assertIsNotNone(m2.table)

    def test_get_unique_cells(self):
        r0_cells = self.mutator.get_unique_cells(0)
        self.assertEqual(len(r0_cells), 2)

        r7_cells = self.mutator.get_unique_cells(7)
        self.assertEqual(len(r7_cells), 4)

        r15_cells = self.mutator.get_unique_cells(15)
        self.assertEqual(len(r15_cells), 4)

        with self.assertRaises(IndexError):
            self.mutator.get_unique_cells(999)

    def test_set_cell_text(self):
        self.mutator.set_cell_text(0, 1, "CÔNG TY CỔ PHẦN NĂNG LƯỢNG MỚI")
        cells = self.mutator.get_unique_cells(0)
        self.assertEqual(cells[1].text, "CÔNG TY CỔ PHẦN NĂNG LƯỢNG MỚI")

        with self.assertRaises(IndexError):
            self.mutator.set_cell_text(0, 99, "Invalid")

    def test_set_prefixed_text(self):
        self.mutator.set_prefixed_text(
            row_index=7,
            cell_index=2,
            prefix="Ngày cấp:",
            value="15/08/2021",
        )
        cells = self.mutator.get_unique_cells(7)
        self.assertEqual(cells[2].text, "Ngày cấp: 15/08/2021")

        self.mutator.set_prefixed_text(
            row_index=13,
            cell_index=1,
            value="500,000",
            suffix="triệu đồng",
        )
        cells13 = self.mutator.get_unique_cells(13)
        self.assertEqual(cells13[1].text, "500,000 triệu đồng")

    def test_set_checkbox(self):
        # Row 3: KH_HIEN_HUU -> index 1
        self.mutator.set_checkbox(3, selected_index=1)
        row3 = self.mutator.table.rows[3]
        ffdata3 = row3._tr.xpath(".//w:ffData")
        self.assertEqual(len(ffdata3), 2)

        # Checkbox 0 unchecked
        cb0 = ffdata3[0].find(qn("w:checkBox"))
        self.assertEqual(cb0.find(qn("w:default")).get(qn("w:val")), "0")
        self.assertEqual(cb0.find(qn("w:checked")).get(qn("w:val")), "0")

        # Checkbox 1 checked
        cb1 = ffdata3[1].find(qn("w:checkBox"))
        self.assertEqual(cb1.find(qn("w:default")).get(qn("w:val")), "1")
        self.assertEqual(cb1.find(qn("w:checked")).get(qn("w:val")), "1")

        # Row 29: TAI_CAP -> index 0
        self.mutator.set_checkbox(29, selected_index=0)
        row29 = self.mutator.table.rows[29]
        ffdata29 = row29._tr.xpath(".//w:ffData")
        cb29_0 = ffdata29[0].find(qn("w:checkBox"))
        cb29_1 = ffdata29[1].find(qn("w:checkBox"))
        self.assertEqual(cb29_0.find(qn("w:checked")).get(qn("w:val")), "1")
        self.assertEqual(cb29_1.find(qn("w:checked")).get(qn("w:val")), "0")

    def test_set_sdt_dropdown_text(self):
        self.mutator.set_sdt_dropdown_text(2, "Công ty TNHH Hai Thành Viên")
        row2 = self.mutator.table.rows[2]
        sdt_content = row2._tr.xpath(".//w:sdtContent")[0]
        t_nodes = sdt_content.findall(".//" + qn("w:t"))
        self.assertEqual(t_nodes[0].text, "Công ty TNHH Hai Thành Viên")

        # Ensure color red is removed
        colors = sdt_content.findall(".//" + qn("w:color"))
        self.assertEqual(len(colors), 0)

    def test_clear_sample_placeholders(self):
        self.mutator.clear_sample_placeholders()

        # Row 0 Cell 1 should be empty
        self.assertEqual(self.mutator.get_unique_cells(0)[1].text, "")

        # Row 15 Cells 1, 2, 3 should be empty
        r15_cells = self.mutator.get_unique_cells(15)
        self.assertEqual(r15_cells[1].text, "")
        self.assertEqual(r15_cells[2].text, "")
        self.assertEqual(r15_cells[3].text, "")

        # Row 16 Cell 2 should be empty
        self.assertEqual(self.mutator.get_unique_cells(16)[2].text, "")

    def test_save_and_reload(self):
        self.mutator.set_cell_text(0, 1, "CONG TY TEST 123")
        self.mutator.set_checkbox(3, 0)
        saved_file = os.path.join(self.tmp_dir, "saved_output.docx")
        self.mutator.save(saved_file)

        # Reload
        reloaded_doc = docx.Document(saved_file)
        reloaded_mutator = SectionADocxMutator(reloaded_doc)
        self.assertEqual(
            reloaded_mutator.get_unique_cells(0)[1].text, "CONG TY TEST 123"
        )
        ffdata = reloaded_mutator.table.rows[3]._tr.xpath(".//w:ffData")
        self.assertEqual(
            ffdata[0].find(qn("w:checkBox")).find(qn("w:checked")).get(qn("w:val")), "1"
        )


if __name__ == "__main__":
    unittest.main()
