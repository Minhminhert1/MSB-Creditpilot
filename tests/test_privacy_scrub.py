# -*- coding: utf-8 -*-
"""Regression tests for the bank-internal-personnel privacy scrub.

Scope reminder (see task): ONLY bank-internal person names (RM, relationship
manager, appraiser, branch director, preparer/checker/approver, etc.) were
removed. Customer-side people (legal representatives, directors, shareholders,
business management named in customer documents) are explicitly OUT OF SCOPE
and must remain fully intact -- these tests assert both directions.
"""

import os
import unittest
import zipfile
from unittest.mock import patch

import docx

import web_copilot_app as w
from msb_eb_copilot.src.document_formatter import DocumentFormatter
from msb_eb_copilot.src.template_verification.verifier import (
    AUTHORITATIVE_MB07_SHA256,
    LEGACY_MB07_SHA256,
    MB07TemplateVerifier,
)

REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
AUTHORITATIVE_TEMPLATE_PATH = os.path.join(
    REPO_ROOT, "MB07 Tờ trình đề xuất cấp tín dụng (ĐVKD) - thực hiện rà soát - tái cấp.docx"
)
LEGACY_TEMPLATE_PATH = os.path.join(
    REPO_ROOT, "MB07 Tờ trình đề xuất cấp tín dụng (ĐVKD) - bản tham khảo.docx"
)

# Confirmed BANK_INTERNAL_PERSON names found during the inventory (RM /
# support / branch-manager fields in CASES_DB, and docProps/comments/people
# metadata embedded in the bundled MB07 templates). Includes the two names the
# task explicitly named as known examples ("Thơ", "Trọng").
CONFIRMED_BANK_STAFF_NAMES = [
    "Lê Văn Hùng",
    "Trần Thị Mai",
    "Vũ Đình Trọng",
    "Phạm Quốc Tuấn",
    "Nguyễn Thị Ngân",
    "Phạm Văn Nam",
    "Trần Tuấn Anh",
    "Hoa Pham Vu Thanh",
    "Phuong Nguyen Hang",
    "Nguyễn Phạm Thanh Ngân",
    "Tho Ngo Thi Ngoc",
    "Ngo Thi Ngoc",
]

# Confirmed CUSTOMER_PERSON names that must NEVER be touched (legal
# representatives, shareholders/board members, customer-side management --
# all sourced from customer/business documentary evidence in CASES_DB).
CONFIRMED_CUSTOMER_NAMES_BY_CASE = {
    "PSD": {
        "legal_rep": "Đại diện Demo",
        "management": ["Phan Hải Âu", "Nguyễn Mạnh Lân", "Lê Minh Kha"],
    },
    "GAS_SOUTH": {
        "legal_rep": "Đặng Văn Vĩnh",
        "management": ["Nguyễn Ngọc Luận", "Đặng Văn Vĩnh", "Trần Thị Thu Thảo"],
    },
    "PHYTOPHARMA": {
        "legal_rep": "Lê Minh Phú",
        "management": ["Nguyễn Công Chiến", "Lê Minh Phú", "Võ Thị Tuấn Anh", "Phan Thị Thu Hà"],
    },
}


def _docx_full_text(doc: "docx.Document") -> str:
    parts = [p.text for p in doc.paragraphs]
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _docx_zip_xml_text(path: str, parts) -> str:
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        return "\n".join(z.read(p).decode("utf-8", errors="ignore") for p in parts if p in names)


class TestBankStaffNamesRemovedFromDemoData(unittest.TestCase):
    """A (demo/preloaded case output) + C/D/E/F (customer-side preserved)."""

    def test_a_no_bank_staff_names_in_cases_db(self):
        cases_db_repr = repr(w.CASES_DB)
        for name in CONFIRMED_BANK_STAFF_NAMES:
            self.assertNotIn(name, cases_db_repr, f"Bank-staff name {name!r} still present in CASES_DB")

    def test_b_bank_role_fields_remain_structurally_present(self):
        """B. rm_metadata keys (RM/support/manager) still exist -- only the
        VALUES were scrubbed to role placeholders, the fields themselves were
        not removed (RM can still fill them in later)."""
        for case_id in ("PSD", "GAS_SOUTH", "PHYTOPHARMA"):
            rm_meta = w.CASES_DB[case_id]["rm_metadata"]
            for key in ("rm_name", "rm_phone", "support_name", "support_phone", "manager_name", "manager_phone"):
                self.assertIn(key, rm_meta, f"{case_id}: missing rm_metadata key {key!r}")
            # Placeholder values are present and non-empty (structurally filled, not deleted).
            self.assertTrue(rm_meta["rm_name"])
            self.assertTrue(rm_meta["support_name"])
            self.assertTrue(rm_meta["manager_name"])
            self.assertIn("tự điền", rm_meta["rm_name"])
            self.assertIn("tự điền", rm_meta["support_name"])
            self.assertIn("tự điền", rm_meta["manager_name"])

    def test_c_customer_legal_representatives_are_not_removed(self):
        for case_id, expected in CONFIRMED_CUSTOMER_NAMES_BY_CASE.items():
            actual = w.CASES_DB[case_id]["customer"]["legal_rep_name"]
            self.assertEqual(actual, expected["legal_rep"], f"{case_id}: legal_rep_name was altered")

    def test_d_legal_representative_names_remain_intact(self):
        # Same guarantee as C, asserted independently per the task's checklist.
        self.assertEqual(w.CASES_DB["GAS_SOUTH"]["customer"]["legal_rep_name"], "Đặng Văn Vĩnh")
        self.assertEqual(w.CASES_DB["PHYTOPHARMA"]["customer"]["legal_rep_name"], "Lê Minh Phú")

    def test_e_shareholder_names_remain_intact(self):
        gas_shareholders = {s["name"] for s in w.CASES_DB["GAS_SOUTH"]["section_c"]["shareholders"]}
        self.assertIn("Tổng Công ty Khí Việt Nam (PV GAS)", gas_shareholders)
        phyto_shareholders = {s["name"] for s in w.CASES_DB["PHYTOPHARMA"]["section_c"]["shareholders"]}
        self.assertIn("Tổng Công ty Dược Việt Nam - CTCP (Vinapharm)", phyto_shareholders)

    def test_f_business_management_names_from_customer_documents_remain_intact(self):
        for case_id, expected in CONFIRMED_CUSTOMER_NAMES_BY_CASE.items():
            actual_names = {m["name"] for m in w.CASES_DB[case_id]["section_c"]["management"]}
            for expected_name in expected["management"]:
                self.assertIn(expected_name, actual_names, f"{case_id}: management name {expected_name!r} missing")


class TestBankStaffNamesRemovedFromBundledTemplates(unittest.TestCase):
    """A (bundled authoritative template) + G (docx still opens) + I (no
    destructive formatting changes)."""

    def test_a_authoritative_template_metadata_contains_no_bank_staff_names(self):
        self.assertTrue(os.path.exists(AUTHORITATIVE_TEMPLATE_PATH))
        xml_text = _docx_zip_xml_text(AUTHORITATIVE_TEMPLATE_PATH, ["docProps/core.xml"])
        for name in CONFIRMED_BANK_STAFF_NAMES:
            self.assertNotIn(name, xml_text, f"{name!r} still present in authoritative template docProps")

    def test_a_legacy_reference_template_metadata_contains_no_bank_staff_names(self):
        """Covers the task's explicit known examples ('Thơ', 'Trọng') -- 'Thơ'
        was embedded as a Word comment author/person identity in this file."""
        self.assertTrue(os.path.exists(LEGACY_TEMPLATE_PATH))
        xml_text = _docx_zip_xml_text(
            LEGACY_TEMPLATE_PATH,
            ["docProps/core.xml", "word/comments.xml", "word/people.xml"],
        )
        for name in CONFIRMED_BANK_STAFF_NAMES:
            self.assertNotIn(name, xml_text, f"{name!r} still present in legacy reference template metadata")

    def test_no_bank_staff_names_in_template_body_text_either(self):
        for path in (AUTHORITATIVE_TEMPLATE_PATH, LEGACY_TEMPLATE_PATH):
            doc = docx.Document(path)
            full_text = _docx_full_text(doc)
            for name in CONFIRMED_BANK_STAFF_NAMES:
                self.assertNotIn(name, full_text, f"{name!r} found in body text of {path}")

    def test_g_both_templates_still_open_successfully(self):
        for path in (AUTHORITATIVE_TEMPLATE_PATH, LEGACY_TEMPLATE_PATH):
            doc = docx.Document(path)  # raises if the package is corrupt
            self.assertGreater(len(doc.paragraphs), 0)
            self.assertGreater(len(doc.tables), 0)

    def test_i_no_destructive_formatting_changes_table_and_paragraph_counts_preserved(self):
        """I. The metadata-only scrub must not have altered document structure --
        pinned against the known-good counts observed before/after the scrub."""
        doc_auth = docx.Document(AUTHORITATIVE_TEMPLATE_PATH)
        self.assertEqual(len(doc_auth.tables), 65)
        self.assertEqual(len(doc_auth.paragraphs), 694)

        doc_legacy = docx.Document(LEGACY_TEMPLATE_PATH)
        self.assertEqual(len(doc_legacy.tables), 73)
        self.assertEqual(len(doc_legacy.paragraphs), 720)

    def test_i_comment_substance_preserved_only_author_identity_removed(self):
        """Only the identifying w:author/w:initials attributes were cleared --
        the review comment's own text content is untouched."""
        xml_text = _docx_zip_xml_text(LEGACY_TEMPLATE_PATH, ["word/comments.xml"])
        self.assertIn("giải thích thêm 2022", xml_text)


class TestTemplateFidelityStillPasses(unittest.TestCase):
    """H. Document QA / template fidelity verification still passes after the
    scrub (the pinned SHA-256 constants were deliberately updated to the new,
    privacy-scrubbed authoritative content)."""

    def test_h_authoritative_template_still_verifies_as_authoritative(self):
        verifier = MB07TemplateVerifier()
        result = verifier.verify(AUTHORITATIVE_TEMPLATE_PATH)
        self.assertTrue(result.is_authoritative)
        self.assertEqual(result.sha256, AUTHORITATIVE_MB07_SHA256)

    def test_h_legacy_template_still_verifies_as_authoritative(self):
        verifier = MB07TemplateVerifier()
        result = verifier.verify(LEGACY_TEMPLATE_PATH)
        self.assertTrue(result.is_authoritative)
        self.assertEqual(result.sha256, LEGACY_MB07_SHA256)


class TestGeneratedMB07OutputNeverLeaksBankStaffNames(unittest.TestCase):
    """A (generated MB07) -- the ROOT-CAUSE, future-proof fix: DocumentFormatter
    now clears author/last-modified-by metadata on every generated proposal,
    regardless of what the template's own docProps says."""

    def test_strip_author_metadata_clears_a_planted_bank_name(self):
        """Directly proves the code-level fix works even if a future template
        re-upload reintroduces a bank employee's name in docProps."""
        import shutil
        import tempfile

        tmp_dir = tempfile.mkdtemp()
        try:
            working_copy = os.path.join(tmp_dir, "working.docx")
            shutil.copy2(AUTHORITATIVE_TEMPLATE_PATH, working_copy)

            doc = docx.Document(working_copy)
            doc.core_properties.author = "Nguyễn Văn Một (Bank Staff Simulated)"
            doc.core_properties.last_modified_by = "Trần Thị Hai (Bank Staff Simulated)"
            doc.save(working_copy)

            output_path = os.path.join(tmp_dir, "polished.docx")
            DocumentFormatter.polish(working_copy, output_path=output_path)

            polished = docx.Document(output_path)
            self.assertEqual(polished.core_properties.author, "")
            self.assertEqual(polished.core_properties.last_modified_by, "")
        finally:
            shutil.rmtree(tmp_dir, ignore_errors=True)

    def test_a_freshly_generated_mb07_output_contains_no_bank_staff_names(self):
        with patch("web_copilot_app.run_document_qa_gate", return_value={"overall_status": "PASS"}):
            output_path = w.execute_generation_pipeline("PSD")

        self.assertTrue(os.path.exists(output_path))
        doc = docx.Document(output_path)

        full_text = _docx_full_text(doc)
        for name in CONFIRMED_BANK_STAFF_NAMES:
            self.assertNotIn(name, full_text, f"{name!r} leaked into generated MB07 body text")

        # Root-cause fix: docProps author/last-modified-by are always cleared.
        self.assertEqual(doc.core_properties.author, "")
        self.assertEqual(doc.core_properties.last_modified_by, "")

    def test_a_freshly_generated_mb07_output_preserves_customer_legal_rep(self):
        """Companion assertion: the SAME generation run must still faithfully
        carry the customer's own legal representative name -- proving the
        scrub did not over-reach into customer-side content."""
        with patch("web_copilot_app.run_document_qa_gate", return_value={"overall_status": "PASS"}):
            output_path = w.execute_generation_pipeline("PSD")
        doc = docx.Document(output_path)
        full_text = _docx_full_text(doc)
        # PSD's legal rep placeholder ("Đại diện Demo") is customer-side and
        # must still be rendered somewhere in the generated proposal.
        self.assertIn("Đại diện Demo", full_text)


if __name__ == "__main__":
    unittest.main()
