# -*- coding: utf-8 -*-
"""Regression tests for the MB07 "Chọn kết quả" placeholder concatenation defect.

Root cause: several cells in the authoritative MB07 template's Section B summary table
("Mục 1", the "Ghi chú" column) hold their default text inside a Word content control
(w:sdt), e.g. "Chọn kết quả". python-docx's Paragraph.runs / Cell.text only ever see
direct <w:r> children of <w:p> and are blind to text nested inside <w:sdt><w:sdtContent>.
SafeCellMutator.set_cell_text used to treat such a cell as having "no runs" and simply
append a brand-new sibling run, leaving the untouched placeholder rendered immediately
before/after the new value -- e.g. "Chọn kết quảTái cấp và ...".

These tests reconstruct text the way Word/LibreOffice actually renders it (all <w:t>
descendants in document order), which is the only way to see this class of defect --
cell.text itself never revealed it, before or after the fix.
"""

from __future__ import annotations

import glob
import os
import shutil
import tempfile
import unittest

import docx

from msb_eb_copilot.src.mb07_inplace_mutator import MB07InPlaceMutator
from msb_eb_copilot.src.template_rendering.safe_mutation import SafeCellMutator


def _full_text(oxml_element) -> str:
    """Visually-rendered text in document order, including text nested inside w:sdt
    content controls -- what Word/LibreOffice renders, unlike python-docx's cell.text."""
    return "".join(t.text or "" for t in oxml_element.xpath(".//w:t"))


class TestMB07PlaceholderConcatenationFix(unittest.TestCase):
    PLACEHOLDER = "Chọn kết quả"

    @classmethod
    def setUpClass(cls):
        base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        candidates = [
            p for p in glob.glob(os.path.join(base_dir, "MB07*.docx")) if "rà soát" in p
        ]
        assert candidates, f"Authoritative MB07 template not found under {base_dir}"
        cls.template_path = candidates[0]

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.test_docx = os.path.join(self.temp_dir, "test_doc.docx")
        shutil.copyfile(self.template_path, self.test_docx)

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def _render(self) -> docx.Document:
        doc = docx.Document(self.test_docx)
        MB07InPlaceMutator(doc).mutate_section_b({})
        doc.save(self.test_docx)
        return docx.Document(self.test_docx)

    def test_template_reproduces_the_defect_shape_before_fix_context(self):
        """Sanity check on the raw template: confirms the exact cells this bug affects
        really do hold their default text inside a w:sdt, invisible to cell.text."""
        doc = docx.Document(self.template_path)
        t5 = doc.tables[5]
        for row_idx in (1, 2, 3, 4):
            cell = t5.rows[row_idx].cells[4]
            self.assertEqual(cell.text, "", msg=f"row {row_idx}")
            self.assertEqual(_full_text(cell._tc), self.PLACEHOLDER, msg=f"row {row_idx}")

    def test_no_populated_cell_concatenates_the_placeholder_with_the_new_value(self):
        doc = self._render()
        t5 = doc.tables[5]

        expected_notes = {
            1: "Tái cấp và nâng HMTD từ 500 tỷ lên 700 tỷ đồng",
            2: "Tối đa 700 tỷ đồng",
            3: "Trong tổng HMTD 700 tỷ đồng",
            4: "Không đề xuất riêng biệt trong kỳ này",
        }
        for row_idx, expected in expected_notes.items():
            cell = t5.rows[row_idx].cells[4]
            full_text = _full_text(cell._tc)
            self.assertEqual(full_text, expected, msg=f"row {row_idx}")
            self.assertNotIn(self.PLACEHOLDER, full_text, msg=f"row {row_idx}")

    def test_every_pure_sdt_placeholder_cell_in_the_document_is_defect_proof(self):
        """Generalizes the fix across the whole template (task requirement: search for
        other occurrences of 'Chọn kết quả' and ensure the same defect cannot happen
        elsewhere). For every cell across the ENTIRE authoritative template whose only
        content is a content-control placeholder -- the exact structural shape of this
        bug: zero direct <w:r> runs, e.g. every "Ghi chú" dropdown cell in table[5], plus
        every other dropdown cell in Section A/D/E -- writing a new value through
        SafeCellMutator must fully replace the placeholder with no concatenation.

        Cells where the sdt sits alongside genuine, non-blank label text (e.g. a fixed
        label ending in ":" immediately followed by its own dropdown) are intentionally
        out of scope here: SafeCellMutator's donor-run path deliberately leaves those
        sdt controls untouched (see test_set_cell_text_with_real_sibling_content_leaves_sdt_untouched),
        since such cells are populated through the dedicated SectionADocxMutator dropdown
        API, never through a blind whole-cell text replacement.
        """
        doc = docx.Document(self.template_path)
        checked = 0
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text != "":
                        continue
                    sdt_t_nodes = cell._tc.xpath(".//w:sdt//w:t")
                    if not sdt_t_nodes:
                        continue
                    checked += 1
                    SafeCellMutator.set_cell_text(cell, "TEST_VALUE_NO_CONCAT")
                    full_text = _full_text(cell._tc)
                    self.assertEqual(full_text, "TEST_VALUE_NO_CONCAT")
                    self.assertNotIn(self.PLACEHOLDER, full_text)

        self.assertGreater(
            checked, 20,
            "Expected to find the template's known population of pure sdt-placeholder cells; "
            "if this drops to 0, the test template may have changed shape and this coverage "
            "guard should be re-examined.",
        )

    def test_content_control_structure_and_styling_preserved(self):
        """The sdt wrapper is preserved (not destructively rebuilt) and the template's
        red/italic placeholder styling is cleared now that real content is final --
        matching the same convention already used for Section A dropdowns."""
        doc = self._render()
        cell = doc.tables[5].rows[1].cells[4]
        tc = cell._tc

        self.assertEqual(len(tc.xpath(".//w:sdt")), 1)
        self.assertEqual(len(cell.paragraphs[0].runs), 0, "no stray sibling run was appended")
        self.assertEqual(len(tc.xpath(".//w:color")), 0)
        self.assertEqual(len(tc.xpath(".//w:i")), 0)

    def test_document_formatter_and_structure_guard_still_pass_after_mutation(self):
        """Guards against a regression where the sdt-handling fix silently violates the
        broader MB07 template fidelity contract (section/table/geometry invariants)."""
        from msb_eb_copilot.src.document_formatter import DocumentFormatter
        from msb_eb_copilot.src.template_rendering.structure_guard import StructureGuard

        self._render()
        DocumentFormatter.polish(self.test_docx, keep_highlights=False)

        guard = StructureGuard(template_path=self.template_path)
        result = guard.verify(self.test_docx, raise_on_violation=True)
        self.assertTrue(result.is_structurally_sound)


if __name__ == "__main__":
    unittest.main(verbosity=2)
