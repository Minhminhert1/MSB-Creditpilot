"""Module: test_anchor_and_safe_mutation.py
Description: Unit and integration tests for MB07 AnchorResolver and Safe Mutation primitives.
"""

import os
import unittest
import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from msb_eb_copilot.src.template_rendering.bindings import (
    MutationType,
    TargetRelationship,
    TemplateBinding,
    MB07_STANDARD_BINDINGS,
)
from msb_eb_copilot.src.template_rendering.anchor_resolver import (
    AnchorResolver,
    AnchorNotFoundError,
    AnchorAmbiguousError,
    TemplateStructureMismatchError,
)
from msb_eb_copilot.src.template_rendering.safe_mutation import (
    SafeCellMutator,
    SafeParagraphMutator,
    DynamicRowCloner,
)


# Mirrors the real MB07 template's dropdown placeholder pattern: a content control
# (w:sdt) whose display run carries the red/italic "Chọn kết quả" default, living
# entirely inside w:sdtContent -- invisible to python-docx's Paragraph.runs.
SDT_PLACEHOLDER_XML = (
    f"<w:sdt {nsdecls('w')}>"
    "<w:sdtPr><w:alias w:val=\"Chọn kết quả\"/></w:sdtPr>"
    "<w:sdtContent>"
    "<w:r><w:rPr><w:color w:val=\"FF0000\"/><w:i/></w:rPr><w:t>Chọn kết quả</w:t></w:r>"
    "</w:sdtContent>"
    "</w:sdt>"
)


def _clear_direct_runs(paragraph) -> None:
    for r in list(paragraph._p.findall(qn("w:r"))):
        paragraph._p.remove(r)


def _inject_sdt_placeholder(paragraph) -> None:
    """Replaces a paragraph's direct runs with a synthetic 'Chọn kết quả' content control,
    reproducing the exact structural shape found in the authoritative MB07 template."""
    _clear_direct_runs(paragraph)
    paragraph._p.append(parse_xml(SDT_PLACEHOLDER_XML))


def _full_text(oxml_element) -> str:
    """Reconstructs the visually-rendered text of an element in document order, including
    text nested inside content controls (w:sdt) -- i.e. what Word/LibreOffice would render,
    as opposed to python-docx's run-only Paragraph.text / Cell.text."""
    return "".join(t.text or "" for t in oxml_element.xpath(".//w:t"))


class TestAnchorAndSafeMutation(unittest.TestCase):

    def setUp(self):
        self.doc = docx.Document()
        # Table 0: Metadata
        t0 = self.doc.add_table(rows=2, cols=2)
        t0.rows[0].cells[0].text = "Đơn vị trình"
        t0.rows[0].cells[1].text = "ORIGINAL_UNIT"
        t0.rows[1].cells[0].text = "Số tờ trình"
        t0.rows[1].cells[1].text = "ORIGINAL_NO"

        # Table 1: Profile
        t1 = self.doc.add_table(rows=2, cols=2)
        t1.rows[0].cells[0].text = "Tên khách hàng"
        t1.rows[0].cells[1].text = "ORIGINAL_NAME"
        t1.rows[1].cells[0].text = "Mã số thuế"
        t1.rows[1].cells[1].text = "ORIGINAL_TAX"

        # Heading & Narrative Paragraphs
        self.doc.add_paragraph("NỘI DUNG ĐỀ XUẤT CẤP TÍN DỤNG")
        t2 = self.doc.add_table(rows=2, cols=3)
        t2.rows[0].cells[0].text = "STT"
        t2.rows[0].cells[1].text = "Nhu cầu"
        t2.rows[0].cells[2].text = "Hạn mức đề xuất"

        p_narr = self.doc.add_paragraph("Đánh giá nguồn trả nợ:")
        r = p_narr.add_run("ORIGINAL_NARRATIVE")
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)

    def test_anchor_find_cell_right_neighbor(self):
        cell = AnchorResolver.find_cell_by_anchor(
            self.doc, "Tên khách hàng", target_rel=TargetRelationship.NEIGHBOR_CELL_RIGHT
        )
        self.assertEqual(cell.text, "ORIGINAL_NAME")

    def test_anchor_not_found_raises(self):
        with self.assertRaises(AnchorNotFoundError):
            AnchorResolver.find_cell_by_anchor(self.doc, "NON_EXISTENT_LABEL")

    def test_anchor_ambiguous_raises(self):
        # Add another cell with same anchor
        t_extra = self.doc.add_table(rows=1, cols=2)
        t_extra.rows[0].cells[0].text = "Tên khách hàng"
        with self.assertRaises(AnchorAmbiguousError):
            AnchorResolver.find_cell_by_anchor(self.doc, "Tên khách hàng")

    def test_find_table_by_heading(self):
        tbl = AnchorResolver.find_table_by_heading(self.doc, "NỘI DUNG ĐỀ XUẤT CẤP TÍN DỤNG")
        self.assertEqual(len(tbl.rows), 2)
        self.assertEqual(tbl.rows[0].cells[1].text, "Nhu cầu")

    def test_safe_cell_mutation_preserves_donor_formatting(self):
        target_cell = self.doc.tables[0].rows[0].cells[1]
        p = target_cell.paragraphs[0]
        r = p.runs[0]
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)
        r.bold = True

        SafeCellMutator.set_cell_text(target_cell, "MUTATED_UNIT", bold=True)

        self.assertEqual(target_cell.text, "MUTATED_UNIT")
        self.assertTrue(target_cell.paragraphs[0].runs[0].bold)
        self.assertEqual(target_cell.paragraphs[0].runs[0].font.name, "Times New Roman")

    def test_safe_cell_multiline_narrative_cloning(self):
        target_cell = self.doc.tables[0].rows[1].cells[1]
        lines = ["Dòng 1: Đánh giá khả năng trả nợ.", "Dòng 2: Nguồn trả nợ khả thi từ dòng tiền."]
        
        SafeCellMutator.set_multiline_text(target_cell, lines, bold=False)

        self.assertEqual(len(target_cell.paragraphs), 2)
        self.assertEqual(target_cell.paragraphs[0].text, lines[0])
        self.assertEqual(target_cell.paragraphs[1].text, lines[1])

    def test_dynamic_row_cloner(self):
        tbl = self.doc.tables[2]
        orig_row_count = len(tbl.rows)

        new_row = DynamicRowCloner.clone_row(tbl, source_row_index=1)

        self.assertEqual(len(tbl.rows), orig_row_count + 1)
        # Verify new row cells are blanked out but present
        self.assertEqual(len(new_row.cells), 3)
        self.assertEqual(new_row.cells[0].text, "")
        
        # Populate new row
        SafeCellMutator.set_cell_text(new_row.cells[0], "1")
        SafeCellMutator.set_cell_text(new_row.cells[1], "Vay hạn mức")
        SafeCellMutator.set_cell_text(new_row.cells[2], "500.000")

        self.assertEqual(new_row.cells[1].text, "Vay hạn mức")

    # ------------------------------------------------------------------
    # Regression: content-control (w:sdt) placeholder must never survive
    # concatenated with the final value (e.g. "Chọn kết quảTái cấp và ...").
    # ------------------------------------------------------------------

    def test_set_cell_text_replaces_sdt_placeholder_without_concatenation(self):
        cell = self.doc.add_table(rows=1, cols=1).rows[0].cells[0]
        _inject_sdt_placeholder(cell.paragraphs[0])

        SafeCellMutator.set_cell_text(cell, "Tái cấp và nâng HMTD từ 500 tỷ lên 700 tỷ đồng")

        full_text = _full_text(cell._tc)
        self.assertEqual(full_text, "Tái cấp và nâng HMTD từ 500 tỷ lên 700 tỷ đồng")
        self.assertNotIn("Chọn kết quả", full_text)
        # The content control wrapper itself is preserved (not destructively rebuilt),
        # and no duplicate sibling run was appended.
        self.assertEqual(len(cell._tc.xpath(".//w:sdt")), 1)
        self.assertEqual(len(cell.paragraphs[0].runs), 0)
        # Placeholder red/italic styling is stripped now that real content is final.
        self.assertEqual(len(cell._tc.xpath(".//w:color")), 0)
        self.assertEqual(len(cell._tc.xpath(".//w:i")), 0)

    def test_set_cell_text_replaces_sdt_placeholder_with_whitespace_sibling_run(self):
        """Some template cells pair the sdt with a purely cosmetic whitespace run; that
        run must not become a decoy donor run that leaves the placeholder untouched."""
        cell = self.doc.add_table(rows=1, cols=1).rows[0].cells[0]
        p = cell.paragraphs[0]
        _inject_sdt_placeholder(p)
        p.add_run("  ")

        SafeCellMutator.set_cell_text(cell, "Tối đa 700 tỷ đồng")

        full_text = _full_text(cell._tc)
        self.assertEqual(full_text, "Tối đa 700 tỷ đồng")
        self.assertNotIn("Chọn kết quả", full_text)

    def test_set_cell_text_with_real_sibling_content_leaves_sdt_untouched(self):
        """If a paragraph already has genuine (non-blank) text alongside an sdt, the
        existing donor-run behavior is preserved and the sdt is left alone."""
        cell = self.doc.add_table(rows=1, cols=1).rows[0].cells[0]
        p = cell.paragraphs[0]
        _inject_sdt_placeholder(p)
        real_run = p.add_run("EXISTING_REAL_CONTENT")

        SafeCellMutator.set_cell_text(cell, "NEW_VALUE")

        self.assertEqual(real_run.text, "NEW_VALUE")

    def test_set_paragraph_text_replaces_sdt_placeholder_without_concatenation(self):
        p = self.doc.add_paragraph()
        _inject_sdt_placeholder(p)

        SafeParagraphMutator.set_paragraph_text(p, "Tối đa 700 tỷ đồng")

        full_text = _full_text(p._p)
        self.assertEqual(full_text, "Tối đa 700 tỷ đồng")
        self.assertNotIn("Chọn kết quả", full_text)
        self.assertEqual(len(p._p.xpath(".//w:sdt")), 1)


if __name__ == "__main__":
    unittest.main()
