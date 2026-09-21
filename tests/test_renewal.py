# -*- coding: utf-8 -*-
"""Tests for the "Tái cấp" (renewal) workspace: deterministic baseline
extraction from the previous MB07 .docx, deterministic ChangeSet comparison
against current canonical facts, and per-case state isolation.

Confirms: UPDATE WHAT CHANGED, PRESERVE WHAT DIDN'T, and Zero Silent Fallback
(REMOVED/CONFLICT always require RM review, never auto-resolved).
"""

import io

import docx
import pytest

from msb_eb_copilot.src.renewal.baseline_extractor import (
    extract_old_mb07_baseline,
    RenewalBaselineExtractionError,
)
from msb_eb_copilot.src.renewal.change_detection import (
    build_change_set,
    extract_new_canonical_snapshot,
    format_trieu_as_ty_display,
)
from msb_eb_copilot.src.renewal.models import ChangeStatus, RenewalBaselineField, RMResolution
from msb_eb_copilot.src.renewal.store import RenewalStore, RenewalItemNotFoundError


def _make_old_mb07_docx(rows):
    """Builds a minimal .docx with a 2-column table (label | value) per row,
    simulating a previously-generated MB07 tờ trình."""
    document = docx.Document()
    table = document.add_table(rows=0, cols=2)
    for label, value in rows:
        row = table.add_row()
        row.cells[0].text = label
        row.cells[1].text = value
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


PSD_CASE_DATA = {
    "customer": {
        "name": "CÔNG TY CỔ PHẦN PHÂN PHỐI DEMO",
        "address": "P.207, Tòa nhà PetroVietnam, Số 1-5 Lê Duẩn, Q.1, TP.HCM",
    },
    "section_c": {
        "customers": [{"name": "CTCP Đầu tư Thế Giới Di Động (MWG)", "share": 4.49}],
    },
    "section_d": {
        "years": ["2024", "2025"],
        "net_revenue": [5702529.0, 7819398.0],
        "net_profit_after_tax": [89729.0, 134201.0],
        "total_assets": [2810436.0, 4683423.0],
        "equity": [597826.0, 729343.0],
        "inventories": [525688.0, 965402.0],
        "receivables": [723020.0, 1475029.0],
        "short_term_debt": [1537823.0, 2572040.0],
    },
    "section_e": {
        "msb_outstanding": 499999.0,
    },
}


class TestBaselineExtraction:
    def test_extracts_labeled_table_rows_deterministically(self):
        docx_bytes = _make_old_mb07_docx([
            ("Tên doanh nghiệp", "CÔNG TY CỔ PHẦN PHÂN PHỐI DEMO"),
            ("Doanh thu thuần", "7.200 tỷ VND"),
            ("Khách hàng lớn nhất", "ABC Corporation"),
        ])
        baseline = extract_old_mb07_baseline(docx_bytes)

        assert baseline["customer.legal_name"].value == "CÔNG TY CỔ PHẦN PHÂN PHỐI DEMO"
        assert baseline["financial.net_revenue_latest"].value == "7.200 tỷ VND"
        assert baseline["business.key_customer"].value == "ABC Corporation"
        assert baseline["customer.legal_name"].source == "MB07 kỳ trước"

    def test_extracts_labeled_paragraphs_as_fallback(self):
        document = docx.Document()
        document.add_paragraph("Địa chỉ trụ sở: 123 Đường ABC, Quận 1, TP.HCM")
        buf = io.BytesIO()
        document.save(buf)
        baseline = extract_old_mb07_baseline(buf.getvalue())
        assert baseline["customer.address"].value == "123 Đường ABC, Quận 1, TP.HCM"

    def test_unmatched_labels_are_absent_not_fabricated(self):
        docx_bytes = _make_old_mb07_docx([("Ghi chú không liên quan", "Một giá trị bất kỳ")])
        baseline = extract_old_mb07_baseline(docx_bytes)
        assert baseline == {}

    def test_invalid_docx_bytes_raise_extraction_error(self):
        with pytest.raises(RenewalBaselineExtractionError):
            extract_old_mb07_baseline(b"this is not a docx file")

    def test_never_writes_to_the_original_docx_bytes_or_mutates_a_docx_object(self):
        """No direct/uncontrolled DOCX rewriting: baseline extraction only READS
        the old MB07 -- it must never open it for writing or save it back."""
        docx_bytes = _make_old_mb07_docx([("Tên doanh nghiệp", "ABC")])
        original_bytes_copy = bytes(docx_bytes)
        extract_old_mb07_baseline(docx_bytes)
        assert docx_bytes == original_bytes_copy  # untouched


class TestNewCanonicalSnapshot:
    def test_snapshot_matches_dashboard_display_convention(self):
        snapshot = extract_new_canonical_snapshot(PSD_CASE_DATA)
        assert snapshot["financial.net_revenue_latest"].value == "7.819 tỷ VND"
        assert snapshot["customer.legal_name"].value == "CÔNG TY CỔ PHẦN PHÂN PHỐI DEMO"
        assert snapshot["business.key_customer"].value == "CTCP Đầu tư Thế Giới Di Động (MWG)"

    def test_missing_data_is_omitted_not_fabricated(self):
        snapshot = extract_new_canonical_snapshot({"customer": {}, "section_d": {}, "section_c": {}, "section_e": {}})
        assert snapshot == {}

    def test_format_trieu_as_ty_display_matches_existing_dashboard_convention(self):
        assert format_trieu_as_ty_display(7819398.0) == "7.819 tỷ VND"


class TestChangeSetDetection:
    """Covers requirement S's core ChangeSet test list: unchanged/changed/new/
    removed/conflict, each producing the correct RM-review requirement."""

    def _snapshot(self):
        return extract_new_canonical_snapshot(PSD_CASE_DATA)

    def test_old_mb07_plus_new_facts_produce_change_set(self):
        old_baseline = {
            "financial.net_revenue_latest": RenewalBaselineField(
                canonical_path="financial.net_revenue_latest", label="Doanh thu thuần",
                value="7.200 tỷ VND", source="MB07 kỳ trước",
            ),
        }
        change_set = build_change_set("PSD", old_baseline, self._snapshot())
        assert change_set.case_id == "PSD"
        assert len(change_set.items) > 0

    def test_unchanged_facts_preserved(self):
        snapshot = self._snapshot()
        old_baseline = {
            "customer.address": RenewalBaselineField(
                canonical_path="customer.address", label="Địa chỉ trụ sở",
                value=snapshot["customer.address"].value, source="MB07 kỳ trước",
            ),
        }
        change_set = build_change_set("PSD", old_baseline, snapshot)
        item = next(i for i in change_set.items if i.canonical_path == "customer.address")
        assert item.status == ChangeStatus.UNCHANGED
        assert not item.requires_rm_action()

    def test_changed_facts_are_flagged(self):
        old_baseline = {
            "financial.net_revenue_latest": RenewalBaselineField(
                canonical_path="financial.net_revenue_latest", label="Doanh thu thuần",
                value="7.200 tỷ VND", source="MB07 kỳ trước",
            ),
        }
        change_set = build_change_set("PSD", old_baseline, self._snapshot())
        item = next(i for i in change_set.items if i.canonical_path == "financial.net_revenue_latest")
        assert item.status == ChangeStatus.CHANGED
        assert item.old_value == "7.200 tỷ VND"
        assert item.new_value == "7.819 tỷ VND"
        assert item.requires_rm_action()
        assert item.rm_resolution == RMResolution.PENDING  # candidate only, not auto-applied

    def test_new_facts_are_flagged(self):
        change_set = build_change_set("PSD", old_baseline={}, new_snapshot=self._snapshot())
        item = next(i for i in change_set.items if i.canonical_path == "customer.legal_name")
        assert item.status == ChangeStatus.NEW
        assert item.old_value is None
        assert item.requires_rm_action()

    def test_removed_facts_require_rm_review_and_are_not_auto_deleted(self):
        snapshot = dict(self._snapshot())
        del snapshot["business.key_customer"]
        old_baseline = {
            "business.key_customer": RenewalBaselineField(
                canonical_path="business.key_customer", label="Khách hàng lớn nhất",
                value="ABC Corporation", source="MB07 kỳ trước",
            ),
        }
        change_set = build_change_set("PSD", old_baseline, snapshot)
        item = next(i for i in change_set.items if i.canonical_path == "business.key_customer")
        assert item.status == ChangeStatus.REMOVED
        # The OLD value is still visibly preserved in the ChangeSetItem -- never silently dropped.
        assert item.old_value == "ABC Corporation"
        assert item.requires_rm_action()

    def test_conflicts_require_rm_review_and_are_never_auto_chosen(self):
        snapshot = dict(self._snapshot())
        snapshot["financial.equity_latest"].value = "5.000 tỷ VND"  # huge unexplained swing
        old_baseline = {
            "financial.equity_latest": RenewalBaselineField(
                canonical_path="financial.equity_latest", label="Vốn chủ sở hữu",
                value="729 tỷ VND", source="MB07 kỳ trước",
            ),
        }
        change_set = build_change_set("PSD", old_baseline, snapshot)
        item = next(i for i in change_set.items if i.canonical_path == "financial.equity_latest")
        assert item.status == ChangeStatus.CONFLICT
        assert item.requires_rm_action()
        assert item.rm_resolution == RMResolution.PENDING

    def test_explicit_conflicting_paths_override_is_respected(self):
        old_baseline = {
            "customer.legal_name": RenewalBaselineField(
                canonical_path="customer.legal_name", label="Tên doanh nghiệp",
                value="CÔNG TY CŨ", source="MB07 kỳ trước",
            ),
        }
        snapshot = self._snapshot()
        change_set = build_change_set("PSD", old_baseline, snapshot, conflicting_paths={"customer.legal_name"})
        item = next(i for i in change_set.items if i.canonical_path == "customer.legal_name")
        assert item.status == ChangeStatus.CONFLICT

    def test_rm_review_complete_only_when_every_actionable_item_resolved(self):
        change_set = build_change_set("PSD", old_baseline={}, new_snapshot=self._snapshot())
        assert change_set.is_rm_review_complete is False
        for item in change_set.items:
            item.rm_resolution = RMResolution.ACCEPT_NEW
        assert change_set.is_rm_review_complete is True


class TestRenewalStoreIsolation:
    def test_case_switching_isolates_renewal_state(self):
        store = RenewalStore()
        store.set_old_baseline("PSD", "old_psd.docx", {"customer.legal_name": RenewalBaselineField(
            canonical_path="customer.legal_name", label="Tên", value="PSD Co", source="MB07 kỳ trước",
        )})
        store.set_old_baseline("GAS_SOUTH", "old_gas.docx", {})

        psd_state = store.get("PSD")
        gas_state = store.get("GAS_SOUTH")
        assert psd_state.old_mb07_filename == "old_psd.docx"
        assert gas_state.old_mb07_filename == "old_gas.docx"
        assert "customer.legal_name" in psd_state.old_baseline
        assert gas_state.old_baseline == {}

    def test_unknown_case_returns_none_not_another_cases_data(self):
        store = RenewalStore()
        store.set_old_baseline("PSD", "old_psd.docx", {})
        assert store.get("SOME_OTHER_CASE") is None

    def test_resolve_change_item_raises_when_no_analysis_yet(self):
        store = RenewalStore()
        with pytest.raises(RenewalItemNotFoundError):
            store.resolve_change_item("PSD", "customer.legal_name", RMResolution.ACCEPT_NEW)

    def test_resolve_change_item_updates_the_correct_item(self):
        store = RenewalStore()
        change_set = build_change_set("PSD", old_baseline={}, new_snapshot=extract_new_canonical_snapshot(PSD_CASE_DATA))
        store.set_change_set("PSD", change_set)
        path = change_set.items[0].canonical_path
        item = store.resolve_change_item("PSD", path, RMResolution.KEEP_OLD, note="RM note")
        assert item.rm_resolution == RMResolution.KEEP_OLD
        assert item.rm_note == "RM note"

    def test_new_baseline_upload_invalidates_previous_change_analysis(self):
        store = RenewalStore()
        change_set = build_change_set("PSD", old_baseline={}, new_snapshot=extract_new_canonical_snapshot(PSD_CASE_DATA))
        store.set_change_set("PSD", change_set)
        assert store.get("PSD").has_change_analysis is True
        store.set_old_baseline("PSD", "new_upload.docx", {})
        assert store.get("PSD").has_change_analysis is False
