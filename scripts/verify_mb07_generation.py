#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Verify MB07 Document Generation Pipeline per Requirement 15."""
import os
import sys
import docx

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

sys.path.insert(0, ".")

from web_copilot_app import execute_generation_pipeline, CASES_DB

def verify():
    print("Testing MB07 Generation Pipeline for PSD...")
    output_path = execute_generation_pipeline("PSD")
    
    # 1. Verify file exists
    assert os.path.exists(output_path), f"File {output_path} does not exist"
    file_size = os.path.getsize(output_path)
    assert file_size > 50000, f"File size too small: {file_size} bytes"
    print(f"[PASS] File generated: {output_path} (size: {file_size:,} bytes)")

    # 2. Verify DOCX can be opened and read by python-docx
    doc = docx.Document(output_path)
    p_count = len(doc.paragraphs)
    t_count = len(doc.tables)
    assert p_count > 10, f"Paragraph count too low: {p_count}"
    assert t_count > 5, f"Table count too low: {t_count}"
    print(f"[PASS] python-docx read successfully: {p_count} paragraphs, {t_count} tables")

    # 3. Verify sections A-E are present in text
    all_text = "\n".join([p.text for p in doc.paragraphs] + [cell.text for t in doc.tables for row in t.rows for cell in row.cells])
    assert "PHẦN A" in all_text or "THÔNG TIN CHUNG" in all_text, "Section A missing"
    assert "PHẦN B" in all_text or "NHU CẦU" in all_text, "Section B missing"
    assert "PHẦN C" in all_text or "HOẠT ĐỘNG KINH DOANH" in all_text, "Section C missing"
    assert "PHẦN D" in all_text or "TÀI CHÍNH" in all_text, "Section D missing"
    assert "PHẦN E" in all_text or "QUAN HỆ TÍN DỤNG" in all_text or "CIC" in all_text, "Section E missing"
    print(f"[PASS] Sections A-E verified in rendered document")

    print("\n" + "=" * 50)
    print("MB07 Generation Pipeline Verification: 100% SUCCESS")
    print("=" * 50)
    return True

if __name__ == "__main__":
    ok = verify()
    sys.exit(0 if ok else 1)
