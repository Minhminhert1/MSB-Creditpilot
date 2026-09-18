# -*- coding: utf-8 -*-
import docx
import sys, os

sys.stdout.reconfigure(encoding='utf-8')
parent_dir = r"c:\Users\minhnh33\Documents\Hackathon"
doc_path = os.path.join(parent_dir, "MB07 Tờ trình đề xuất cấp tín dụng (ĐVKD) - Tho.docx")
doc = docx.Document(doc_path)

print("=== TABLE 0 (HEADER BOX) ===")
t0 = doc.tables[0]
for r in t0.rows:
    print([c.text.replace('\n', ' ').strip() for c in r.cells])

print("\n=== TABLE 1 (CUSTOMER PROFILE) ===")
t1 = doc.tables[1]
for r in t1.rows[:12]:
    print([c.text.replace('\n', ' ').strip() for c in r.cells])

print("\n=== TABLE 2 (LIMIT PROPOSAL) ===")
t2 = doc.tables[2]
for r in t2.rows:
    print([c.text.replace('\n', ' ').strip() for c in r.cells])
