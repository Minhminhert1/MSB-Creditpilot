# -*- coding: utf-8 -*-
import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

import os

parent_dir = r"c:\Users\minhnh33\Documents\Hackathon"
files = [
    os.path.join(parent_dir, "Tờ trình mẫu mới - Gas South.docx"),
    os.path.join(parent_dir, "MB07 Tờ trình đề xuất cấp tín dụng (ĐVKD) - Tho.docx"),
    os.path.join(parent_dir, "THÉP TÂY ĐÔ_ Tờ trình_2026.docx")
]

for f in files:
    try:
        doc = docx.Document(f)
        print(f"\n{'='*60}\nFILE: {f}\n{'='*60}")
        print(f"Total Paragraphs: {len(doc.paragraphs)}, Total Tables: {len(doc.tables)}")
        
        print("\n--- HEADINGS & MAIN SECTIONS ---")
        for i, p in enumerate(doc.paragraphs):
            t = p.text.strip()
            if any(t.startswith(x) for x in ["PHẦN", "MỤC", "Phần", "Mục", "I.", "II.", "III.", "IV.", "V.", "VI.", "VII.", "1.", "2.", "3.", "A.", "B.", "C."]) and len(t) < 120:
                print(f"P{i:03d} | {t}")
                
        print("\n--- SAMPLE TABLES ---")
        for idx, tbl in enumerate(doc.tables[:6]):
            print(f"Table {idx}: {len(tbl.rows)} rows x {len(tbl.columns)} cols")
            if len(tbl.rows) > 0:
                h = [c.text.replace('\n', ' ').strip()[:35] for c in tbl.rows[0].cells]
                print(f"   Header: {h}")
    except Exception as e:
        print(f"Error reading {f}: {e}")
