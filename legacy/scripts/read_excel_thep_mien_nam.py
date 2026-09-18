# -*- coding: utf-8 -*-
import openpyxl, os, sys
if hasattr(sys.stdout, "reconfigure"): sys.stdout.reconfigure(encoding="utf-8")

folder = r"c:\Users\minhnh33\Documents\Hackathon\New folder"
excel_path = os.path.join(folder, "MB09.QT.RR.044 Template BCTC 2025.updated T2.2206 - Copy.xlsx")
wb = openpyxl.load_workbook(excel_path, data_only=True)

print("=== CHECKING SHEET 04.TỜ TRÌNH FOR REVENUE & EQUITY ===")
ws_tt = wb["04.Tờ trình"]
for r in ws_tt.iter_rows(values_only=True):
    txt = str(r[0])
    if any(k in txt.lower() for k in ["doanh thu", "vốn", "lợi nhuận", "tổng tài sản"]):
        print([c for c in r if c is not None][:6])

print("\n=== CHECKING SHEET 05. CIC ===")
ws_cic = wb["05. CIC"]
for r in ws_cic.iter_rows(values_only=True):
    if any(r):
        print([str(c)[:30] for c in r if c is not None][:6])
